from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


@dataclass
class ResumeParseResult:
    filename: str
    file_type: str
    text: str
    word_count: int
    page_count: int | None


class ResumeParserError(Exception):
    """Raised when an uploaded resume cannot be parsed."""


def clean_extracted_text(text: str) -> str:
    """
    Normalize whitespace while preserving readable line breaks.
    """
    text = text.replace("\x00", " ")
    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    cleaned_lines: list[str] = []

    for line in text.splitlines():
        cleaned_line = re.sub(
            r"[ \t]+",
            " ",
            line,
        ).strip()

        if cleaned_line:
            cleaned_lines.append(cleaned_line)

    return "\n".join(cleaned_lines)


def extract_pdf_text(
    file_bytes: bytes,
) -> tuple[str, int]:
    try:
        document = pymupdf.open(
            stream=file_bytes,
            filetype="pdf",
        )
    except Exception as exc:
        raise ResumeParserError(
            "The PDF could not be opened."
        ) from exc

    page_text: list[str] = []

    try:
        for page in document:
            extracted_text = page.get_text(
                "text",
                sort=True,
            )

            if extracted_text.strip():
                page_text.append(extracted_text)

        page_count = document.page_count

    finally:
        document.close()

    combined_text = "\n".join(page_text)

    return combined_text, page_count


def extract_docx_text(
    file_bytes: bytes,
) -> str:
    try:
        document = Document(
            io.BytesIO(file_bytes)
        )
    except Exception as exc:
        raise ResumeParserError(
            "The DOCX file could not be opened."
        ) from exc

    extracted_parts: list[str] = []

    for paragraph in document.paragraphs:
        paragraph_text = paragraph.text.strip()

        if paragraph_text:
            extracted_parts.append(
                paragraph_text
            )

    # Some resumes place important content inside tables.
    for table in document.tables:
        for row in table.rows:
            row_values = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if row_values:
                extracted_parts.append(
                    " | ".join(row_values)
                )

    return "\n".join(extracted_parts)


def parse_resume(
    filename: str,
    file_bytes: bytes,
) -> ResumeParseResult:
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ResumeParserError(
            "Only PDF and DOCX resumes are supported."
        )

    if not file_bytes:
        raise ResumeParserError(
            "The uploaded resume is empty."
        )

    page_count: int | None = None

    if extension == ".pdf":
        raw_text, page_count = extract_pdf_text(
            file_bytes
        )

    else:
        raw_text = extract_docx_text(
            file_bytes
        )

    cleaned_text = clean_extracted_text(
        raw_text
    )

    if not cleaned_text:
        raise ResumeParserError(
            "No readable text was found. "
            "The resume may be scanned or image-based."
        )

    return ResumeParseResult(
        filename=filename,
        file_type=extension.removeprefix(
            "."
        ).upper(),
        text=cleaned_text,
        word_count=len(
            cleaned_text.split()
        ),
        page_count=page_count,
    )
