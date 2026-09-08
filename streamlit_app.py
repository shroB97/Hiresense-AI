"""HireSense AI - self-contained, evidence-grounded resume matcher.

Run with:
    pip install streamlit pandas pypdf python-docx reportlab openai
    streamlit run streamlit_app.py

The OpenAI dependency is optional. Without an API key the application uses a
deterministic matcher. With a key it uses structured AI extraction and
assessment, then applies the same evidence and specificity guardrails.
"""

from __future__ import annotations

import csv
import html
import io
import json
import os
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date
from enum import Enum
from typing import Any, Iterable, Sequence
from urllib.parse import quote_plus

import streamlit as st


st.set_page_config(page_title="HireSense AI", page_icon="◎", layout="wide")


# ---------------------------------------------------------------------------
# Models and configuration
# ---------------------------------------------------------------------------


class Category(str, Enum):
    EDUCATION = "education"
    EXPERIENCE = "experience"
    RESPONSIBILITIES = "responsibilities"
    TECHNICAL_SKILLS = "technical_skills"
    TOOLS = "tools"
    LEADERSHIP = "leadership"
    DOMAIN_KNOWLEDGE = "domain_knowledge"
    CERTIFICATIONS = "certifications"
    SOFT_SKILLS = "soft_skills"
    METHODOLOGIES = "methodologies"
    OTHER = "other"


class Importance(str, Enum):
    REQUIRED = "required"
    PREFERRED = "preferred"


class MatchStatus(str, Enum):
    DIRECT = "direct_match"
    RELATED = "related_evidence"
    UNCERTAIN = "uncertain"
    MISSING = "not_found"


STATUS_CREDIT: dict[MatchStatus, float] = {
    MatchStatus.DIRECT: 1.00,
    MatchStatus.RELATED: 0.65,
    MatchStatus.UNCERTAIN: 0.25,
    MatchStatus.MISSING: 0.00,
}

IMPORTANCE_WEIGHT: dict[Importance, float] = {
    Importance.REQUIRED: 2.0,
    Importance.PREFERRED: 1.0,
}


@dataclass(frozen=True)
class Settings:
    api_key: str = ""
    model: str = "gpt-4.1-mini"
    max_upload_mb: int = 10
    min_resume_chars: int = 80
    use_ai: bool = True


@dataclass(frozen=True)
class ParsedResume:
    filename: str
    text: str
    file_type: str
    page_count: int
    word_count: int


@dataclass(frozen=True)
class Evidence:
    evidence_id: str
    text: str
    source: str
    section: str


@dataclass(frozen=True)
class CandidateProfile:
    candidate_name: str
    professional_summary: str
    total_experience_years: float | None
    skills: tuple[str, ...]
    tools: tuple[str, ...]
    education: tuple[str, ...]
    certifications: tuple[str, ...]
    evidence: tuple[Evidence, ...]
    raw_text: str


@dataclass(frozen=True)
class Requirement:
    requirement_id: str
    text: str
    category: Category
    importance: Importance


@dataclass(frozen=True)
class RequirementMatch:
    requirement: Requirement
    status: MatchStatus
    confidence: float
    assessment: str
    evidence: Evidence | None = None


@dataclass(frozen=True)
class CategoryScore:
    category: Category
    percentage: int
    requirements: int
    direct: int
    related: int
    uncertain: int
    missing: int


@dataclass(frozen=True)
class MatchReport:
    score: int
    alignment: str
    matches: tuple[RequirementMatch, ...]
    category_scores: tuple[CategoryScore, ...]
    required_score: int | None
    preferred_score: int | None
    mode: str
    warnings: tuple[str, ...] = field(default_factory=tuple)


@dataclass(frozen=True)
class ImprovementRecommendation:
    requirement_id: str
    requirement: str
    priority: str
    action: str
    rationale: str
    potential_points: float
    requires_verification: bool
    current_evidence: str = ""


@dataclass(frozen=True)
class BulletRewrite:
    evidence_id: str
    original: str
    improved: str
    reason: str
    requirement_ids: tuple[str, ...] = field(default_factory=tuple)


@dataclass(frozen=True)
class OptimizedResume:
    headline: str
    professional_summary: str
    skills: tuple[str, ...]
    rewrites: tuple[BulletRewrite, ...]
    mode: str
    warnings: tuple[str, ...] = field(default_factory=tuple)


class ResumeParserError(ValueError):
    pass


def _secret(name: str, default: str = "") -> str:
    try:
        value = st.secrets.get(name, default)
    except Exception:
        value = default
    return str(value or default)


def get_settings(session_api_key: str = "", model: str = "") -> Settings:
    key = session_api_key.strip() or _secret("OPENAI_API_KEY") or os.getenv("OPENAI_API_KEY", "")
    selected_model = model.strip() or _secret("OPENAI_MODEL", "gpt-4.1-mini")
    return Settings(api_key=key, model=selected_model, use_ai=bool(key))


# ---------------------------------------------------------------------------
# Text normalization and parsing
# ---------------------------------------------------------------------------


SECTION_NAMES: dict[str, str] = {
    "summary": "Professional Summary",
    "professional summary": "Professional Summary",
    "profile": "Professional Summary",
    "professional profile": "Professional Summary",
    "career summary": "Professional Summary",
    "experience": "Experience",
    "professional experience": "Experience",
    "work experience": "Experience",
    "employment history": "Experience",
    "career history": "Experience",
    "skills": "Skills",
    "technical skills": "Skills",
    "core skills": "Skills",
    "core competencies": "Skills",
    "competencies": "Skills",
    "tools": "Tools",
    "technologies": "Tools",
    "technical proficiencies": "Tools",
    "education": "Education",
    "academic background": "Education",
    "academic qualifications": "Education",
    "certifications": "Certifications",
    "certification": "Certifications",
    "licenses and certifications": "Certifications",
    "projects": "Projects",
    "project experience": "Projects",
    "awards": "Awards",
    "achievements": "Achievements",
    "publications": "Publications",
}

BULLET_RE = re.compile(r"^\s*[\u2022\u25cf\u25aa\u25e6\u2043\u2219*\-–—▪]+\s*")
SPACE_RE = re.compile(r"\s+")
DATE_RANGE_RE = re.compile(
    r"\b(?:19|20)\d{2}\b\s*(?:-|–|—|to)\s*(?:\b(?:19|20)\d{2}\b|present|current|now)",
    re.I,
)


def clean_line(value: str) -> str:
    value = value.replace("\u00a0", " ").replace("\uf0b7", "•")
    value = BULLET_RE.sub("", value)
    return SPACE_RE.sub(" ", value).strip()


def normalize(value: str) -> str:
    value = value.casefold()
    value = value.replace("&", " and ")
    value = re.sub(r"(?<=\w)[/](?=\w)", " ", value)
    value = re.sub(r"[^a-z0-9+#.]+", " ", value)
    return SPACE_RE.sub(" ", value).strip()


def heading_name(line: str) -> str | None:
    cleaned = clean_line(line).rstrip(":").strip()
    key = normalize(cleaned)
    if key in SECTION_NAMES:
        return SECTION_NAMES[key]
    # Do not treat every short all-caps line as a section: names and tools such
    # as "SHROBANTI BANERJEE" and "SQL" are commonly capitalized in resumes.
    return None


def _dedupe(items: Iterable[str]) -> list[str]:
    output: list[str] = []
    seen: set[str] = set()
    for item in items:
        cleaned = clean_line(item)
        key = normalize(cleaned)
        if not cleaned or key in seen:
            continue
        seen.add(key)
        output.append(cleaned)
    return output


def _extract_pdf(data: bytes) -> tuple[str, int]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ResumeParserError("PDF support requires pypdf: pip install pypdf") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ResumeParserError("The PDF is password-protected and cannot be read.") from exc
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
    except ResumeParserError:
        raise
    except Exception as exc:
        raise ResumeParserError(f"The PDF could not be read: {exc}") from exc
    return "\n".join(page for page in pages if page), len(reader.pages)


def _extract_docx(data: bytes) -> tuple[str, int]:
    try:
        from docx import Document
    except ImportError as exc:
        raise ResumeParserError("DOCX support requires python-docx: pip install python-docx") from exc
    try:
        document = Document(io.BytesIO(data))
        blocks: list[str] = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = _dedupe(cell.text for cell in row.cells)
                if cells:
                    blocks.append(" | ".join(cells))
    except Exception as exc:
        raise ResumeParserError(f"The DOCX file could not be read: {exc}") from exc
    return "\n".join(blocks), 0


def parse_resume(filename: str, data: bytes, mime_type: str, settings: Settings) -> ParsedResume:
    if not data:
        raise ResumeParserError("The uploaded resume is empty.")
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise ResumeParserError(f"Upload a resume smaller than {settings.max_upload_mb} MB.")
    suffix = os.path.splitext(filename.casefold())[1]
    if suffix == ".pdf" or mime_type == "application/pdf":
        text, pages = _extract_pdf(data)
        file_type = "PDF"
    elif suffix == ".docx" or "wordprocessingml" in (mime_type or ""):
        text, pages = _extract_docx(data)
        file_type = "DOCX"
    else:
        raise ResumeParserError("Use a PDF or DOCX resume.")
    text = "\n".join(line.rstrip() for line in text.replace("\r", "\n").splitlines())
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(re.sub(r"\s+", "", text)) < settings.min_resume_chars:
        raise ResumeParserError(
            "Very little text was extracted. If this is a scanned PDF, export it with OCR or upload DOCX."
        )
    return ParsedResume(filename, text, file_type, pages, len(text.split()))


def sectionize_resume(text: str) -> list[tuple[str, list[str]]]:
    sections: list[tuple[str, list[str]]] = []
    current_name = "Header"
    current_lines: list[str] = []
    for raw in text.splitlines():
        if not raw.strip():
            continue
        detected = heading_name(raw)
        if detected:
            if current_lines:
                sections.append((current_name, current_lines))
            current_name, current_lines = detected, []
        else:
            current_lines.append(raw.strip())
    if current_lines:
        sections.append((current_name, current_lines))
    return sections


def _is_contact_line(line: str) -> bool:
    return bool(
        re.search(r"\b[\w.+-]+@[\w.-]+\.[a-z]{2,}\b", line, re.I)
        or re.search(r"(?:https?://|linkedin\.com|github\.com)", line, re.I)
        or re.fullmatch(r"[+()\d\s.-]{7,}", line)
    )


def _looks_like_context(line: str) -> bool:
    cleaned = clean_line(line)
    if DATE_RANGE_RE.search(cleaned):
        return True
    if len(cleaned) > 110 or len(cleaned.split()) > 14:
        return False
    if re.search(r"[.!?]$", cleaned):
        return False
    return bool(re.search(r"\s[|–—]\s|\bat\b", cleaned, re.I))


def _is_resume_context_line(line: str) -> bool:
    if _looks_like_context(line) or DATE_RANGE_RE.search(line):
        return True
    cleaned = clean_line(line)
    if len(cleaned.split()) > 10 or re.search(r"[.!?;]$", cleaned):
        return False
    action_verb = re.match(
        r"^(achieved|analyzed|automated|built|conducted|created|delivered|designed|developed|drove|"
        r"facilitated|implemented|improved|led|managed|migrated|optimized|partnered|performed|reduced|"
        r"standardized|standardised|supported|transformed|used|utilized)\b",
        normalize(cleaned),
    )
    return not bool(action_verb)


def split_list_items(lines: Sequence[str]) -> list[str]:
    joined = "\n".join(lines)
    parts = re.split(r"\n|\s*[|;•]\s*|,(?=\s*[A-Z0-9])", joined)
    return _dedupe(parts)


def build_candidate_profile(parsed: ParsedResume) -> CandidateProfile:
    sections = sectionize_resume(parsed.text)
    evidence_rows: list[tuple[str, str, str]] = []
    skills: list[str] = []
    tools: list[str] = []
    education: list[str] = []
    certifications: list[str] = []
    summary = ""

    header_lines = next((lines for name, lines in sections if name == "Header"), [])
    name = next(
        (
            clean_line(line)
            for line in header_lines
            if not _is_contact_line(line)
            and 1 < len(clean_line(line)) < 70
            and len(clean_line(line).split()) <= 6
        ),
        "Candidate",
    )

    for section, lines in sections:
        cleaned_lines = _dedupe(lines)
        if not cleaned_lines:
            continue
        if section == "Professional Summary":
            summary = " ".join(cleaned_lines)
            evidence_rows.append((summary, section, section))
            continue
        if section == "Skills":
            skills.extend(split_list_items(cleaned_lines))
        elif section == "Tools":
            tools.extend(split_list_items(cleaned_lines))
        elif section == "Education":
            education.extend(cleaned_lines)
        elif section == "Certifications":
            certifications.extend(cleaned_lines)

        if section in {"Experience", "Projects"}:
            context: list[str] = []
            for raw, cleaned in ((raw, clean_line(raw)) for raw in lines):
                if not cleaned or _is_contact_line(cleaned):
                    continue
                was_bullet = bool(BULLET_RE.match(raw))
                if not was_bullet and _looks_like_context(cleaned):
                    if not DATE_RANGE_RE.fullmatch(cleaned):
                        context_item = re.sub(DATE_RANGE_RE, "", cleaned).strip(" |–—,-")
                        context_item = re.sub(r"\s*\|\s*", " — ", context_item).strip(" —")
                        context.append(context_item)
                        context = [item for item in context[-2:] if item]
                    continue
                source = " — ".join(context) if context else section
                evidence_rows.append((cleaned, source, section))
        else:
            for cleaned in cleaned_lines:
                if not _is_contact_line(cleaned):
                    evidence_rows.append((cleaned, section, section))

    if not summary:
        candidate_summary = [
            clean_line(line)
            for line in header_lines
            if not _is_contact_line(line) and clean_line(line) != name and len(clean_line(line)) > 50
        ]
        summary = " ".join(candidate_summary[:3])
        if summary:
            evidence_rows.insert(0, (summary, "Professional Summary", "Professional Summary"))

    # Include concise list items independently so a named skill/tool can be retrieved exactly.
    for item in _dedupe(skills):
        evidence_rows.append((item, "Skills", "Skills"))
    for item in _dedupe(tools):
        evidence_rows.append((item, "Tools", "Tools"))

    unique_rows: list[tuple[str, str, str]] = []
    seen_rows: set[tuple[str, str]] = set()
    for text, source, section in evidence_rows:
        key = (normalize(text), normalize(source))
        if len(text) < 2 or key in seen_rows:
            continue
        seen_rows.add(key)
        unique_rows.append((text, source, section))
    evidence = tuple(
        Evidence(f"E{index:04d}", text, source, section)
        for index, (text, source, section) in enumerate(unique_rows, start=1)
    )

    years = extract_total_experience_years(parsed.text)
    return CandidateProfile(
        candidate_name=name,
        professional_summary=summary,
        total_experience_years=years,
        skills=tuple(_dedupe(skills)),
        tools=tuple(_dedupe(tools)),
        education=tuple(_dedupe(education)),
        certifications=tuple(_dedupe(certifications)),
        evidence=evidence,
        raw_text=parsed.text,
    )


def extract_total_experience_years(text: str) -> float | None:
    explicit: list[float] = []
    for match in re.finditer(
        r"\b(\d{1,2}(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)\b(?:\s+of\s+(?:professional\s+)?)?",
        text,
        re.I,
    ):
        value = float(match.group(1))
        if 0 < value <= 70:
            explicit.append(value)
    if explicit:
        return max(explicit)

    ranges: list[tuple[int, int]] = []
    current_year = date.today().year
    for match in DATE_RANGE_RE.finditer(text):
        years = re.findall(r"\b(?:19|20)\d{2}\b", match.group(0))
        if not years:
            continue
        start = int(years[0])
        end = int(years[1]) if len(years) > 1 else current_year
        if start <= end <= current_year + 1:
            ranges.append((start, end))
    if not ranges:
        return None
    # Merge overlapping employment ranges so concurrent roles are not double-counted.
    ranges.sort()
    merged: list[list[int]] = []
    for start, end in ranges:
        if not merged or start > merged[-1][1]:
            merged.append([start, end])
        else:
            merged[-1][1] = max(merged[-1][1], end)
    years = sum(end - start for start, end in merged)
    return float(years) if years > 0 else None


# ---------------------------------------------------------------------------
# Job requirement extraction
# ---------------------------------------------------------------------------


JOB_SECTION_HEADINGS = {
    "required qualifications": (Importance.REQUIRED, "qualification"),
    "minimum qualifications": (Importance.REQUIRED, "qualification"),
    "basic qualifications": (Importance.REQUIRED, "qualification"),
    "requirements": (Importance.REQUIRED, "qualification"),
    "qualifications": (Importance.REQUIRED, "qualification"),
    "what you need": (Importance.REQUIRED, "qualification"),
    "must have": (Importance.REQUIRED, "qualification"),
    "preferred qualifications": (Importance.PREFERRED, "qualification"),
    "desired qualifications": (Importance.PREFERRED, "qualification"),
    "nice to have": (Importance.PREFERRED, "qualification"),
    "preferred": (Importance.PREFERRED, "qualification"),
    "responsibilities": (Importance.REQUIRED, "responsibility"),
    "key responsibilities": (Importance.REQUIRED, "responsibility"),
    "what you will do": (Importance.REQUIRED, "responsibility"),
    "what you'll do": (Importance.REQUIRED, "responsibility"),
    "the role": (Importance.REQUIRED, "responsibility"),
    "duties": (Importance.REQUIRED, "responsibility"),
}

KNOWN_TOOLS = (
    "SAP Signavio",
    "Signavio",
    "SAP S/4HANA",
    "SAP ERP",
    "Oracle",
    "JDE",
    "JD Edwards",
    "ARIS",
    "iGrafx",
    "Bizagi",
    "Camunda",
    "Visio",
    "Power BI",
    "Tableau",
    "Looker",
    "Excel",
    "PowerPoint",
    "Word",
    "SharePoint",
    "Microsoft 365",
    "Power Automate",
    "Python",
    "SQL",
    "Celonis",
    "Alteryx",
    "Snowflake",
    "AWS",
    "Azure",
    "GCP",
    "Jira",
    "Confluence",
    "ServiceNow",
    "Salesforce",
)

VALID_CATEGORIES = {item.value: item for item in Category}
VALID_IMPORTANCE = {item.value: item for item in Importance}


def _job_heading(line: str) -> tuple[Importance, str] | None:
    key = normalize(clean_line(line).rstrip(":"))
    if key in JOB_SECTION_HEADINGS:
        return JOB_SECTION_HEADINGS[key]
    return None


def categorize_requirement(text: str, section_type: str = "") -> Category:
    value = normalize(text)
    if re.search(r"\b(degree|bachelor|master|mba|phd|education)\b", value):
        return Category.EDUCATION
    if re.search(r"\b(certification|certified|green belt|black belt|license)\b", value):
        return Category.CERTIFICATIONS
    if re.search(r"\b(years?|yrs?)\b.*\bexperience\b|\bexperience\b.*\b(years?|yrs?)\b", value):
        return Category.EXPERIENCE
    if any(normalize(tool) in value for tool in KNOWN_TOOLS):
        return Category.TOOLS
    if re.search(r"\b(bpmn|apqc|pcf|raci|agile|scrum|lean|six sigma|methodolog|taxonomy|metadata)\b", value):
        return Category.METHODOLOGIES
    if re.search(r"\b(process architecture|process controls?|domain knowledge|business acumen|industry knowledge)\b", value):
        return Category.DOMAIN_KNOWLEDGE
    if re.search(r"\b(influence|mentor|leadership|change agent|manage a team|people manager)\b", value):
        return Category.LEADERSHIP
    if re.search(r"\b(sql|python|data analysis|analytics|statistical|machine learning|technical skill)\b", value):
        return Category.TECHNICAL_SKILLS
    if section_type != "responsibility" and re.search(
        r"\b(collaboration|collaborative|communication|facilitation|facilitate|analytical capabilities|interpersonal|problem.solving)\b",
        value,
    ):
        return Category.SOFT_SKILLS
    if section_type == "responsibility" or re.match(
        r"^(support|maintain|document|ensure|drive|facilitate|utilize|use|develop|track|prepare|monitor|analyze|assist|act|partner|manage|lead|create|provide|perform|coordinate|identify)\b",
        value,
    ):
        return Category.RESPONSIBILITIES
    if "experience" in value:
        return Category.EXPERIENCE
    return Category.OTHER


def _split_job_line(line: str) -> list[str]:
    cleaned = clean_line(line).strip(" ;")
    if not cleaned:
        return []
    # Office applications are independently testable requirements.
    office = [tool for tool in ("Excel", "PowerPoint", "Word", "SharePoint") if re.search(rf"\b{tool}\b", cleaned, re.I)]
    if len(office) >= 2:
        prefix = "Strong Microsoft Office skills:"
        return [f"{prefix} {tool}" for tool in office]
    semicolon_parts = [part.strip() for part in re.split(r"\s*;\s*", cleaned) if part.strip()]
    if len(semicolon_parts) > 1:
        return semicolon_parts
    if len(cleaned) > 220:
        sentence_parts = [
            part.strip(" .")
            for part in re.split(r"(?<=[.!?])\s+(?=[A-Z])", cleaned)
            if len(part.strip()) >= 12
        ]
        if len(sentence_parts) > 1:
            return sentence_parts
    return [cleaned.rstrip(".")]


def deterministic_requirements(job_description: str) -> list[Requirement]:
    raw_lines = [line for line in job_description.replace("\r", "\n").splitlines()]
    importance = Importance.REQUIRED
    section_type = ""
    candidates: list[tuple[str, Importance, str]] = []
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        paragraph = " ".join(clean_line(part) for part in paragraph_buffer)
        for item in _split_job_line(paragraph):
            candidates.append((item, importance, section_type))
        paragraph_buffer = []

    for raw in raw_lines:
        stripped = raw.strip()
        if not stripped:
            flush_paragraph()
            continue
        detected = _job_heading(stripped)
        if detected:
            flush_paragraph()
            importance, section_type = detected
            continue
        is_bullet = bool(BULLET_RE.match(raw))
        cleaned = clean_line(raw)
        if is_bullet:
            flush_paragraph()
            for item in _split_job_line(cleaned):
                candidates.append((item, importance, section_type))
        else:
            # A short colon-terminated label is probably an unrecognized heading.
            if cleaned.endswith(":") and len(cleaned.split()) <= 7:
                flush_paragraph()
                lowered = normalize(cleaned)
                if any(word in lowered for word in ("preferred", "desired", "nice")):
                    importance = Importance.PREFERRED
                elif any(word in lowered for word in ("required", "minimum", "must")):
                    importance = Importance.REQUIRED
                section_type = "responsibility" if "responsib" in lowered or "dut" in lowered else section_type
                continue
            paragraph_buffer.append(cleaned)
    flush_paragraph()

    # If pasted text has no bullets/newlines, split it into sentences.
    if len(candidates) <= 2:
        candidates = []
        for sentence in re.split(r"(?<=[.!?])\s+(?=[A-Z])", SPACE_RE.sub(" ", job_description)):
            sentence = clean_line(sentence).strip(" .")
            if len(sentence.split()) >= 3:
                inferred_importance = (
                    Importance.PREFERRED
                    if re.search(r"\b(preferred|desired|nice to have|a plus)\b", sentence, re.I)
                    else Importance.REQUIRED
                )
                for item in _split_job_line(sentence):
                    candidates.append((item, inferred_importance, ""))

    boilerplate = re.compile(
        r"\b(equal opportunity|all qualified applicants|accommodation|privacy policy|benefits include|salary range|compensation range|apply now|about us|our company)\b",
        re.I,
    )
    output: list[Requirement] = []
    seen: set[str] = set()
    for text, item_importance, item_section in candidates:
        text = clean_line(text)
        key = normalize(text)
        if len(text) < 12 or len(text.split()) < 2 or boilerplate.search(text) or key in seen:
            continue
        if len(text) > 420:
            text = text[:417].rstrip() + "..."
            key = normalize(text)
        seen.add(key)
        category = categorize_requirement(text, item_section)
        output.append(Requirement(f"R{len(output)+1:03d}", text, category, item_importance))
    if not output:
        raise ValueError("No job requirements could be extracted. Paste the responsibilities and qualifications sections.")
    return output[:100]


def _json_object(raw: str) -> dict[str, Any]:
    value = raw.strip()
    if value.startswith("```"):
        value = re.sub(r"^```(?:json)?\s*|\s*```$", "", value, flags=re.I)
    start, end = value.find("{"), value.rfind("}")
    if start < 0 or end < start:
        raise ValueError("The AI response did not contain a JSON object.")
    parsed = json.loads(value[start : end + 1])
    if not isinstance(parsed, dict):
        raise ValueError("The AI response was not a JSON object.")
    return parsed


def _chat_json(settings: Settings, system: str, user: str) -> dict[str, Any]:
    try:
        from openai import OpenAI
    except ImportError as exc:
        raise RuntimeError("AI mode requires the openai package: pip install openai") from exc
    client = OpenAI(api_key=settings.api_key, timeout=90.0, max_retries=2)
    response = client.chat.completions.create(
        model=settings.model,
        temperature=0,
        response_format={"type": "json_object"},
        messages=[
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
    )
    content = response.choices[0].message.content
    if not content:
        raise RuntimeError("The AI response was empty.")
    return _json_object(content)


def ai_requirements(job_description: str, settings: Settings) -> list[Requirement]:
    system = """You extract atomic, independently evaluable requirements from job descriptions.
Return JSON only: {"requirements":[{"text":"...","category":"...","importance":"required|preferred"}]}.
Allowed categories: education, experience, responsibilities, technical_skills, tools, leadership,
domain_knowledge, certifications, soft_skills, methodologies, other.

Rules:
- Include each real responsibility, required qualification, and preferred qualification once.
- Preserve the employer's meaning; remove promotional, EEO, compensation, and benefits text.
- A bullet containing one coherent responsibility stays together.
- Split separately testable named Office applications (Excel, PowerPoint, Word, SharePoint).
- Keep alternatives together (for example SAP, Oracle, JDE, or comparable platforms).
- "Required" includes day-to-day responsibilities and minimum/basic qualifications.
- "Preferred" applies only when the job explicitly marks it preferred, desired, optional, or a plus.
- Do not invent requirements, tools, years, certifications, or frameworks."""
    payload = _chat_json(settings, system, "JOB DESCRIPTION:\n" + job_description[:30000])
    rows = payload.get("requirements")
    if not isinstance(rows, list):
        raise ValueError("AI extraction returned no requirements list.")
    output: list[Requirement] = []
    seen: set[str] = set()
    for row in rows:
        if not isinstance(row, dict):
            continue
        text = clean_line(str(row.get("text", ""))).rstrip(".")
        key = normalize(text)
        if len(text) < 8 or key in seen:
            continue
        seen.add(key)
        category = VALID_CATEGORIES.get(str(row.get("category", "")).casefold(), categorize_requirement(text))
        importance = VALID_IMPORTANCE.get(str(row.get("importance", "")).casefold(), Importance.REQUIRED)
        output.append(Requirement(f"R{len(output)+1:03d}", text, category, importance))
    if not output:
        raise ValueError("AI extraction did not return usable requirements.")
    return output[:100]


def extract_requirements(job_description: str, settings: Settings, use_ai: bool) -> tuple[list[Requirement], str, list[str]]:
    warnings: list[str] = []
    if use_ai and settings.api_key:
        try:
            return ai_requirements(job_description, settings), "Structured AI", warnings
        except Exception as exc:
            warnings.append(f"AI requirement extraction was unavailable; deterministic extraction was used ({exc}).")
    return deterministic_requirements(job_description), "Deterministic", warnings


# ---------------------------------------------------------------------------
# Evidence retrieval and matching
# ---------------------------------------------------------------------------


STOPWORDS = {
    "a", "an", "and", "are", "as", "at", "be", "by", "for", "from", "have", "in", "into",
    "is", "it", "of", "on", "or", "our", "that", "the", "their", "this", "to", "using", "with",
    "within", "work", "working", "ability", "strong", "knowledge", "understanding", "experience",
    "skills", "skill", "support", "including", "such", "other", "related", "role", "day", "responsibility",
}

CONCEPT_GROUPS = (
    {"document", "documenting", "documentation", "documented", "map", "mapping", "model", "modeling", "flow", "flows", "workflow", "workflows", "bpmn"},
    {"governance", "govern", "standard", "standards", "standardization", "standardised", "standardized", "repository", "auditability", "compliance"},
    {"process", "processes", "business-process", "operational"},
    {"transform", "transformation", "redesign", "improvement", "improve", "optimization", "optimisation", "continuous"},
    {"stakeholder", "stakeholders", "cross-functional", "collaboration", "partnered", "workshop", "workshops", "facilitate", "facilitated"},
    {"metric", "metrics", "kpi", "kpis", "scorecard", "scorecards", "dashboard", "dashboards", "performance", "analytics", "reporting"},
    {"change", "adoption", "enablement", "implementation", "implement", "implemented", "uat"},
    {"requirement", "requirements", "feedback", "decision", "decisions", "action", "actions", "recommendation", "recommendations"},
    {"risk", "risks", "control", "controls", "audit", "auditability", "compliance"},
    {"global", "enterprise", "shared-services", "ssc"},
    {"procure-to-pay", "p2p", "accounts-payable", "ap"},
    {"order-to-cash", "otc"},
    {"record-to-report", "rtr"},
)

ALIASES: dict[str, tuple[str, ...]] = {
    "sap signavio": ("signavio",),
    "sap s/4hana": ("s4hana", "s 4hana"),
    "jd edwards": ("jde",),
    "power bi": ("powerbi",),
    "power automate": ("microsoft power automate",),
    "microsoft 365": ("office 365", "m365"),
    "bpmn 2.0": ("bpmn",),
    "lean six sigma": ("six sigma",),
    "shared services": ("shared service", "ssc"),
    "global process owner": ("gpo",),
    "procure to pay": ("p2p", "procure-to-pay"),
    "order to cash": ("otc", "order-to-cash"),
    "record to report": ("rtr", "record-to-report"),
}


def canonical_text(value: str) -> str:
    result = f" {normalize(value)} "
    for canonical, variants in ALIASES.items():
        options = (canonical,) + variants
        if any(f" {normalize(option)} " in result for option in options):
            result += " " + " ".join(normalize(option) for option in options)
    return SPACE_RE.sub(" ", result).strip()


def concept_tokens(value: str) -> set[str]:
    base = {token for token in canonical_text(value).split() if token not in STOPWORDS and len(token) > 1}
    expanded = set(base)
    for group in CONCEPT_GROUPS:
        if base.intersection(group):
            expanded.update(group)
    return expanded


def phrase_present(phrase: str, text: str) -> bool:
    phrase_norm, text_norm = normalize(phrase), canonical_text(text)
    options = (phrase,) + ALIASES.get(phrase_norm, ())
    return any(re.search(rf"(?<!\w){re.escape(normalize(option))}(?!\w)", text_norm) for option in options)


def lexical_score(requirement: Requirement, evidence: Evidence) -> float:
    req_tokens = concept_tokens(requirement.text)
    ev_tokens = concept_tokens(evidence.text)
    if not req_tokens or not ev_tokens:
        return 0.0
    overlap = len(req_tokens & ev_tokens)
    coverage = overlap / max(1, min(len(req_tokens), 12))
    precision = overlap / max(1, min(len(ev_tokens), 18))
    score = 0.72 * coverage + 0.28 * precision
    req_norm, ev_norm = canonical_text(requirement.text), canonical_text(evidence.text)
    for tool in KNOWN_TOOLS:
        if phrase_present(tool, req_norm) and phrase_present(tool, ev_norm):
            score += 0.34
    section = evidence.section.casefold()
    expected_sections = {
        Category.EDUCATION: {"education"},
        Category.CERTIFICATIONS: {"certifications"},
        Category.TOOLS: {"tools", "skills", "experience", "projects"},
        Category.EXPERIENCE: {"experience", "professional summary"},
        Category.RESPONSIBILITIES: {"experience", "projects"},
    }.get(requirement.category, {"experience", "projects", "skills", "professional summary"})
    if section in expected_sections:
        score += 0.08
    if normalize(evidence.text) and normalize(evidence.text) in req_norm and len(evidence.text) >= 3:
        score += 0.18
    return min(score, 1.0)


def retrieve_evidence(requirement: Requirement, evidence: Sequence[Evidence], limit: int = 8) -> list[tuple[Evidence, float]]:
    ranked = sorted(
        ((item, lexical_score(requirement, item)) for item in evidence),
        key=lambda pair: (pair[1], len(pair[0].text)),
        reverse=True,
    )
    useful = [pair for pair in ranked if pair[1] >= 0.05]
    return (useful or ranked)[:limit]


def _named_tools(text: str) -> list[str]:
    return [tool for tool in KNOWN_TOOLS if phrase_present(tool, text)]


def _required_years(text: str) -> float | None:
    matches = re.findall(r"\b(\d{1,2}(?:\.\d+)?)\s*(?:-|–|—|to)?\s*\d*\s*\+?\s*(?:years?|yrs?)\b", text, re.I)
    return float(matches[0]) if matches else None


def _degree_level(text: str) -> str | None:
    value = normalize(text)
    if re.search(r"\b(phd|doctorate|doctoral)\b", value):
        return "doctorate"
    if re.search(r"\b(master|masters|msc|m\.s\.?|mba)\b", value):
        return "master"
    if re.search(r"\b(bachelor|bachelors|btech|b\.s\.?|bsc|undergraduate degree)\b", value):
        return "bachelor"
    return None


def specificity_guard(
    requirement: Requirement,
    evidence: Evidence | None,
    proposed: MatchStatus,
    candidate: CandidateProfile,
) -> MatchStatus:
    """Downgrade claims that require an exact credential, named framework, or named tool."""
    if evidence is None:
        return MatchStatus.MISSING if proposed in {MatchStatus.DIRECT, MatchStatus.RELATED} else proposed
    req, ev = canonical_text(requirement.text), canonical_text(evidence.text)

    required_years = _required_years(req)
    if required_years is not None and proposed == MatchStatus.DIRECT:
        if candidate.total_experience_years is None:
            return MatchStatus.UNCERTAIN
        if candidate.total_experience_years + 0.01 < required_years:
            return MatchStatus.RELATED if candidate.total_experience_years >= required_years * 0.7 else MatchStatus.MISSING

    if requirement.category == Category.EDUCATION and proposed == MatchStatus.DIRECT:
        required_level, actual_level = _degree_level(req), _degree_level(ev)
        order = {"bachelor": 1, "master": 2, "doctorate": 3}
        if required_level and (not actual_level or order[actual_level] < order[required_level]):
            return MatchStatus.RELATED if actual_level else MatchStatus.UNCERTAIN

    if requirement.category == Category.CERTIFICATIONS:
        asks_signavio = "signavio" in req and ("certification" in req or "certified" in req)
        asks_belt = bool(re.search(r"\b(green belt|black belt)\b", req))
        explicit_credential = evidence.section.casefold() == "certifications" or bool(
            re.search(r"\b(certification|certified|credential|license)\b", ev)
        )
        exact_credential = (
            asks_signavio and "signavio" in ev and explicit_credential
        ) or (asks_belt and bool(re.search(r"\b(green belt|black belt)\b", ev)))
        if (asks_signavio or asks_belt) and proposed == MatchStatus.DIRECT and not exact_credential:
            return MatchStatus.RELATED
        if proposed == MatchStatus.DIRECT and not explicit_credential:
            return MatchStatus.RELATED

    exact_only_terms = ("apqc", "raci", "sharepoint", "taxonomy", "metadata")
    for term in exact_only_terms:
        if re.search(rf"\b{term}\b", req) and proposed == MatchStatus.DIRECT and not re.search(rf"\b{term}\b", ev):
            return MatchStatus.RELATED

    tools = _named_tools(req)
    if tools and requirement.category == Category.TOOLS and proposed == MatchStatus.DIRECT:
        exact = [tool for tool in tools if phrase_present(tool, ev)]
        # "SAP/Oracle/JDE or comparable" is directly met by any explicitly named alternative.
        if not exact:
            return MatchStatus.RELATED
        if "sap erp" in req and "sap signavio" in ev and "sap erp" not in ev:
            return MatchStatus.RELATED

    return proposed


def deterministic_match(requirement: Requirement, candidate: CandidateProfile) -> RequirementMatch:
    retrieved = retrieve_evidence(requirement, candidate.evidence)
    evidence, score = retrieved[0] if retrieved else (None, 0.0)
    req_norm = canonical_text(requirement.text)
    ev_norm = canonical_text(evidence.text) if evidence else ""

    status = MatchStatus.MISSING
    if evidence:
        tools = _named_tools(req_norm)
        exact_tool = any(phrase_present(tool, ev_norm) for tool in tools)
        req_tokens = concept_tokens(req_norm)
        ev_tokens = concept_tokens(ev_norm)
        coverage = len(req_tokens & ev_tokens) / max(1, min(len(req_tokens), 12))
        if exact_tool or score >= 0.62 or (coverage >= 0.52 and len(req_tokens & ev_tokens) >= 2):
            status = MatchStatus.DIRECT
        elif score >= 0.30 or (coverage >= 0.28 and len(req_tokens & ev_tokens) >= 2):
            status = MatchStatus.RELATED
        elif score >= 0.16:
            status = MatchStatus.UNCERTAIN

    if requirement.category == Category.EDUCATION and evidence:
        required_level, actual_level = _degree_level(req_norm), _degree_level(ev_norm)
        order = {"bachelor": 1, "master": 2, "doctorate": 3}
        if required_level and actual_level and order[actual_level] >= order[required_level]:
            status = MatchStatus.DIRECT

    required_years = _required_years(req_norm)
    if required_years is not None:
        if candidate.total_experience_years is None:
            status = min(status, MatchStatus.UNCERTAIN, key=lambda item: STATUS_CREDIT[item])
        elif candidate.total_experience_years >= required_years and evidence:
            status = MatchStatus.DIRECT if score >= 0.22 else MatchStatus.RELATED
        elif candidate.total_experience_years >= required_years * 0.7:
            status = MatchStatus.RELATED
        else:
            status = MatchStatus.MISSING

    status = specificity_guard(requirement, evidence, status, candidate)
    if status == MatchStatus.DIRECT:
        confidence = min(0.97, max(0.84, 0.80 + score * 0.17))
        assessment = "The résumé explicitly demonstrates this requirement through the cited evidence."
    elif status == MatchStatus.RELATED:
        confidence = min(0.92, max(0.68, 0.62 + score * 0.30))
        assessment = "The résumé shows closely related capability, but it does not explicitly verify every part of this requirement."
    elif status == MatchStatus.UNCERTAIN:
        confidence = min(0.75, max(0.50, 0.46 + score * 0.35))
        assessment = "The résumé contains a possible supporting signal, but the evidence is too ambiguous for match credit."
    else:
        confidence = max(0.75, min(0.96, 0.91 - score * 0.30))
        assessment = "No sufficiently specific supporting evidence was found in the résumé."
        evidence = None
    return RequirementMatch(requirement, status, confidence, assessment, evidence)


def ai_match_batch(
    requirements: Sequence[Requirement],
    candidate: CandidateProfile,
    settings: Settings,
) -> list[RequirementMatch]:
    evidence_lookup = {item.evidence_id: item for item in candidate.evidence}
    request_rows: list[dict[str, Any]] = []
    retrieval_scores: dict[str, dict[str, float]] = {}
    for requirement in requirements:
        candidates = retrieve_evidence(requirement, candidate.evidence)
        retrieval_scores[requirement.requirement_id] = {
            evidence.evidence_id: score for evidence, score in candidates
        }
        request_rows.append(
            {
                "requirement_id": requirement.requirement_id,
                "requirement": requirement.text,
                "category": requirement.category.value,
                "importance": requirement.importance.value,
                "resume_candidates": [
                    {
                        "evidence_id": evidence.evidence_id,
                        "source": evidence.source,
                        "text": evidence.text,
                    }
                    for evidence, _score in candidates
                ],
            }
        )
    system = """You are a strict, evidence-grounded resume-to-job assessor.
Return JSON only: {"matches":[{"requirement_id":"R001","status":"direct_match|related_evidence|uncertain|not_found","confidence":0.0,"evidence_id":"E0001 or null","assessment":"one or two sentences"}]}.

Definitions:
- direct_match: the selected resume excerpt explicitly proves the full requirement or a clearly accepted alternative.
- related_evidence: the excerpt proves an adjacent/transferable capability, but an explicit term, scope, context, or credential is absent.
- uncertain: a weak or ambiguous signal that should receive very limited credit.
- not_found: no supplied excerpt substantively supports the requirement.

Mandatory safeguards:
- Evaluate every requirement independently. Do not let overall seniority inflate a specific item.
- Use only one supplied evidence_id for each assessment. Never quote, paraphrase, or invent resume facts outside it.
- A named certification is direct only when that exact certification/credential is explicit.
- A named framework (APQC, RACI), platform, or tool is direct only when explicitly named, except where the requirement itself accepts comparable alternatives.
- General Microsoft 365 experience is related, not direct, for an individual app that is not named.
- Signavio is not proof of core SAP ERP experience.
- Do not infer undocumented controls, risks, decision points, ownership structures, metadata, or taxonomy work.
- Leadership verbs can prove influence only when the excerpt describes cross-functional leadership or stakeholder outcomes.
- If status is not_found, evidence_id must be null. Keep confidence calibrated; do not default everything to 95%."""
    payload = _chat_json(
        settings,
        system,
        json.dumps(
            {
                "candidate_total_experience_years": candidate.total_experience_years,
                "requirements": request_rows,
            },
            ensure_ascii=False,
        ),
    )
    rows = payload.get("matches")
    if not isinstance(rows, list):
        raise ValueError("AI assessment returned no matches list.")
    row_lookup = {
        str(row.get("requirement_id")): row
        for row in rows
        if isinstance(row, dict) and row.get("requirement_id")
    }
    results: list[RequirementMatch] = []
    for requirement in requirements:
        row = row_lookup.get(requirement.requirement_id)
        if not row:
            results.append(deterministic_match(requirement, candidate))
            continue
        try:
            status = MatchStatus(str(row.get("status", "not_found")))
        except ValueError:
            status = MatchStatus.MISSING
        evidence_id = str(row.get("evidence_id") or "")
        evidence = evidence_lookup.get(evidence_id)
        # The model may select only evidence that retrieval supplied for this requirement.
        if evidence_id not in retrieval_scores.get(requirement.requirement_id, {}):
            evidence = None
        status = specificity_guard(requirement, evidence, status, candidate)
        if status == MatchStatus.MISSING:
            evidence = None
        try:
            confidence = float(row.get("confidence", 0.5))
        except (TypeError, ValueError):
            confidence = 0.5
        confidence = min(0.99, max(0.05, confidence))
        assessment = clean_line(str(row.get("assessment", "")))
        if not assessment:
            assessment = deterministic_match(requirement, candidate).assessment
        results.append(RequirementMatch(requirement, status, confidence, assessment, evidence))
    return results


def assess_requirements(
    requirements: Sequence[Requirement],
    candidate: CandidateProfile,
    settings: Settings,
    use_ai: bool,
) -> tuple[list[RequirementMatch], str, list[str]]:
    warnings: list[str] = []
    if use_ai and settings.api_key:
        completed: list[RequirementMatch] = []
        try:
            for start in range(0, len(requirements), 10):
                completed.extend(ai_match_batch(requirements[start : start + 10], candidate, settings))
            return completed, "Structured AI + evidence guardrails", warnings
        except Exception as exc:
            warnings.append(f"AI evidence assessment was unavailable; deterministic assessment was used ({exc}).")
    return [deterministic_match(item, candidate) for item in requirements], "Deterministic evidence matcher", warnings


def weighted_percentage(matches: Sequence[RequirementMatch]) -> int | None:
    if not matches:
        return None
    numerator = sum(
        STATUS_CREDIT[item.status] * IMPORTANCE_WEIGHT[item.requirement.importance]
        for item in matches
    )
    denominator = sum(IMPORTANCE_WEIGHT[item.requirement.importance] for item in matches)
    return round(100 * numerator / denominator) if denominator else None


def alignment_label(score: int) -> str:
    if score >= 85:
        return "Excellent alignment"
    if score >= 70:
        return "Strong alignment"
    if score >= 55:
        return "Moderate alignment"
    if score >= 40:
        return "Partial alignment"
    return "Limited alignment"


def create_match_report(
    candidate: CandidateProfile,
    requirements: Sequence[Requirement],
    settings: Settings,
    use_ai: bool,
    initial_warnings: Sequence[str] = (),
    extraction_mode: str = "",
) -> MatchReport:
    matches, matching_mode, matching_warnings = assess_requirements(requirements, candidate, settings, use_ai)
    score = weighted_percentage(matches) or 0
    by_category: dict[Category, list[RequirementMatch]] = defaultdict(list)
    for item in matches:
        by_category[item.requirement.category].append(item)
    category_scores: list[CategoryScore] = []
    for category in Category:
        rows = by_category.get(category, [])
        if not rows:
            continue
        counts = Counter(item.status for item in rows)
        category_scores.append(
            CategoryScore(
                category=category,
                percentage=weighted_percentage(rows) or 0,
                requirements=len(rows),
                direct=counts[MatchStatus.DIRECT],
                related=counts[MatchStatus.RELATED],
                uncertain=counts[MatchStatus.UNCERTAIN],
                missing=counts[MatchStatus.MISSING],
            )
        )
    required = [item for item in matches if item.requirement.importance == Importance.REQUIRED]
    preferred = [item for item in matches if item.requirement.importance == Importance.PREFERRED]
    mode = f"{extraction_mode} extraction · {matching_mode}" if extraction_mode else matching_mode
    return MatchReport(
        score=score,
        alignment=alignment_label(score),
        matches=tuple(matches),
        category_scores=tuple(category_scores),
        required_score=weighted_percentage(required),
        preferred_score=weighted_percentage(preferred),
        mode=mode,
        warnings=tuple(initial_warnings) + tuple(matching_warnings),
    )


# ---------------------------------------------------------------------------
# Score-improvement plan and evidence-safe resume optimization
# ---------------------------------------------------------------------------


def _score_denominator(report: MatchReport) -> float:
    return sum(IMPORTANCE_WEIGHT[item.requirement.importance] for item in report.matches)


def build_improvement_recommendations(report: MatchReport) -> list[ImprovementRecommendation]:
    denominator = _score_denominator(report) or 1.0
    recommendations: list[ImprovementRecommendation] = []
    for item in report.matches:
        if item.status == MatchStatus.DIRECT:
            continue
        required = item.requirement.importance == Importance.REQUIRED
        if required and item.status in {MatchStatus.MISSING, MatchStatus.UNCERTAIN}:
            priority = "Critical"
        elif required:
            priority = "High"
        elif item.status in {MatchStatus.MISSING, MatchStatus.UNCERTAIN}:
            priority = "Medium"
        else:
            priority = "Low"
        potential = (
            100
            * (STATUS_CREDIT[MatchStatus.DIRECT] - STATUS_CREDIT[item.status])
            * IMPORTANCE_WEIGHT[item.requirement.importance]
            / denominator
        )
        if item.status == MatchStatus.RELATED and item.evidence:
            action = (
                "Clarify the scope, context, and outcome in the cited résumé bullet using the employer's terminology—"
                "but only where that wording is factually supported."
            )
            rationale = "The résumé already contains related evidence, so clearer phrasing may make the match explicit."
            requires_verification = False
        elif item.status == MatchStatus.UNCERTAIN:
            action = (
                "Verify whether you performed this work. If true, add one concrete example with the action, process/tool, "
                "stakeholders, and measurable or observable outcome."
            )
            rationale = "The current signal is ambiguous and cannot safely be presented as experience yet."
            requires_verification = True
        else:
            action = (
                "Do not add this merely as a keyword. Add it only if true, supported by a project or role, and explain "
                "how you applied it."
            )
            rationale = "No sufficiently specific résumé evidence currently supports this requirement."
            requires_verification = True
        recommendations.append(
            ImprovementRecommendation(
                requirement_id=item.requirement.requirement_id,
                requirement=item.requirement.text,
                priority=priority,
                action=action,
                rationale=rationale,
                potential_points=potential,
                requires_verification=requires_verification,
                current_evidence=item.evidence.text if item.evidence else "",
            )
        )
    rank = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3}
    return sorted(recommendations, key=lambda item: (rank[item.priority], -item.potential_points, item.requirement_id))


def related_clarity_ceiling(report: MatchReport) -> int:
    """A conservative what-if score: related rows become direct; unsupported rows stay unchanged."""
    denominator = _score_denominator(report)
    if not denominator:
        return report.score
    numerator = 0.0
    for item in report.matches:
        credit = 1.0 if item.status == MatchStatus.RELATED and item.evidence else STATUS_CREDIT[item.status]
        numerator += credit * IMPORTANCE_WEIGHT[item.requirement.importance]
    return round(100 * numerator / denominator)


def _existing_headline(candidate: CandidateProfile) -> str:
    sections = sectionize_resume(candidate.raw_text)
    header = next((lines for name, lines in sections if name == "Header"), [])
    for line in header:
        cleaned = clean_line(line)
        if cleaned == candidate.candidate_name or _is_contact_line(cleaned):
            continue
        if 2 <= len(cleaned.split()) <= 12 and len(cleaned) <= 100:
            return cleaned
    return ""


def supported_keywords(candidate: CandidateProfile, report: MatchReport) -> list[str]:
    """Return only keywords already supported by direct/related resume evidence."""
    ranked: list[tuple[int, str]] = []
    for item in report.matches:
        if item.status not in {MatchStatus.DIRECT, MatchStatus.RELATED} or not item.evidence:
            continue
        base_priority = 0 if item.requirement.importance == Importance.REQUIRED else 2
        if item.status == MatchStatus.RELATED:
            base_priority += 1
        for tool in _named_tools(item.requirement.text):
            if phrase_present(tool, candidate.raw_text):
                ranked.append((base_priority, tool))
        for phrase in (
            "BPMN",
            "process mapping",
            "process governance",
            "process transformation",
            "stakeholder management",
            "change management",
            "business requirements",
            "continuous improvement",
            "operational excellence",
            "KPI reporting",
            "process analytics",
        ):
            if phrase_present(phrase, item.requirement.text) and lexical_score(
                item.requirement, item.evidence
            ) >= 0.22:
                ranked.append((base_priority + 1, phrase.title() if phrase != "BPMN" else phrase))
    for item in candidate.skills + candidate.tools:
        ranked.append((3, item))
    output: list[str] = []
    seen: set[str] = set()
    for _priority, keyword in sorted(ranked, key=lambda pair: pair[0]):
        key = normalize(keyword)
        if not key or key in seen:
            continue
        seen.add(key)
        output.append(keyword)
    return output[:30]


def _specific_unsupported_phrases(candidate: CandidateProfile, report: MatchReport) -> set[str]:
    phrases: set[str] = set()
    raw = canonical_text(candidate.raw_text)
    for item in report.matches:
        for tool in _named_tools(item.requirement.text):
            if not phrase_present(tool, raw):
                phrases.add(tool)
        req = canonical_text(item.requirement.text)
        for phrase in (
            "APQC",
            "RACI",
            "taxonomy",
            "metadata",
            "Green Belt",
            "Black Belt",
            "SAP Signavio certification",
        ):
            if normalize(phrase) in req and normalize(phrase) not in raw:
                phrases.add(phrase)
    return phrases


def _contains_forbidden(value: str, forbidden: Iterable[str]) -> bool:
    normalized = canonical_text(value)
    return any(normalize(phrase) in normalized for phrase in forbidden)


def _number_tokens(value: str) -> set[str]:
    return set(re.findall(r"(?<!\w)\d+(?:\.\d+)?(?:%|\+)?", value))


def _numbers_are_grounded(rewrite: str, source: str) -> bool:
    return _number_tokens(rewrite).issubset(_number_tokens(source))


def _rewrite_is_grounded(rewrite: str, source: str) -> bool:
    if _number_tokens(rewrite) != _number_tokens(source):
        return False
    source_tools = {normalize(tool) for tool in _named_tools(source)}
    rewrite_tools = {normalize(tool) for tool in _named_tools(rewrite)}
    return rewrite_tools.issubset(source_tools)


def deterministic_optimized_resume(
    candidate: CandidateProfile,
    report: MatchReport,
    job_title: str,
    warning: str = "",
) -> OptimizedResume:
    headline = _existing_headline(candidate) or (f"Target Role: {job_title}" if job_title else "Professional Profile")
    warnings = (warning,) if warning else ()
    return OptimizedResume(
        headline=headline,
        professional_summary=candidate.professional_summary,
        skills=tuple(supported_keywords(candidate, report)),
        rewrites=(),
        mode="Conservative formatting and supported-keyword prioritization",
        warnings=warnings,
    )


def ai_optimized_resume(
    candidate: CandidateProfile,
    report: MatchReport,
    job_title: str,
    company: str,
    settings: Settings,
) -> OptimizedResume:
    eligible: dict[str, dict[str, Any]] = {}
    requirement_map: dict[str, list[dict[str, str]]] = defaultdict(list)
    for match in report.matches:
        if match.status not in {MatchStatus.DIRECT, MatchStatus.RELATED} or not match.evidence:
            continue
        if match.evidence.section.casefold() not in {"experience", "projects"}:
            continue
        eligible[match.evidence.evidence_id] = {
            "evidence_id": match.evidence.evidence_id,
            "source": match.evidence.source,
            "original": match.evidence.text,
        }
        requirement_map[match.evidence.evidence_id].append(
            {
                "requirement_id": match.requirement.requirement_id,
                "requirement": match.requirement.text,
                "status": match.status.value,
            }
        )
    evidence_rows: list[dict[str, Any]] = []
    for evidence_id, row in list(eligible.items())[:35]:
        evidence_rows.append({**row, "supported_requirements": requirement_map[evidence_id]})
    allowed_skills = supported_keywords(candidate, report)
    system = """You are a truthful executive resume editor. Return JSON only:
{"headline":"...","professional_summary":"...","skills":["..."],"rewrites":[{"evidence_id":"E0001","improved":"...","reason":"...","requirement_ids":["R001"]}]}.

Your objective is recruiter readability and ATS alignment, never score manipulation through fabrication.
Rules:
- Use only facts in the supplied resume evidence. Do not invent tools, certifications, frameworks, employers, titles,
  dates, responsibilities, scope, metrics, leadership, clients, industries, or outcomes.
- Preserve every number exactly. Do not calculate or introduce a new number.
- A rewrite must stay semantically equivalent to its one original evidence excerpt. It may improve action verbs,
  clarity, concision, and placement of supported job terminology.
- Do not convert related evidence into an unsupported claim. Do not add any missing/uncertain requirement.
- Use only supplied evidence_id values, at most one rewrite per evidence_id.
- Select skills only from allowed_skills and copy their spelling exactly.
- The summary must be 55-90 words, grounded in the supplied evidence, and may use only supported requirements.
- The headline must describe the existing profile; do not claim a certification or job title the candidate has not held.
- requirement_ids must be selected only from that evidence row's supported_requirements.
- Omit weak rewrites. Quality is more important than rewriting every bullet."""
    payload = _chat_json(
        settings,
        system,
        json.dumps(
            {
                "target_job_title": job_title,
                "target_company": company,
                "existing_headline": _existing_headline(candidate),
                "existing_summary": candidate.professional_summary,
                "allowed_skills": allowed_skills,
                "resume_evidence": evidence_rows,
            },
            ensure_ascii=False,
        ),
    )
    forbidden = _specific_unsupported_phrases(candidate, report)
    headline = clean_line(str(payload.get("headline", ""))) or _existing_headline(candidate)
    if not headline or _contains_forbidden(headline, forbidden) or not _numbers_are_grounded(headline, candidate.raw_text):
        headline = _existing_headline(candidate) or (f"Target Role: {job_title}" if job_title else "Professional Profile")
    summary = clean_line(str(payload.get("professional_summary", "")))
    if (
        not summary
        or not 35 <= len(summary.split()) <= 120
        or _contains_forbidden(summary, forbidden)
        or not _numbers_are_grounded(summary, candidate.raw_text)
    ):
        summary = candidate.professional_summary

    allowed_lookup = {normalize(item): item for item in allowed_skills}
    selected_skills: list[str] = []
    for item in payload.get("skills", []):
        selected = allowed_lookup.get(normalize(str(item)))
        if selected and normalize(selected) not in {normalize(value) for value in selected_skills}:
            selected_skills.append(selected)
    for item in allowed_skills:
        if len(selected_skills) >= 30:
            break
        if normalize(item) not in {normalize(value) for value in selected_skills}:
            selected_skills.append(item)

    rewrites: list[BulletRewrite] = []
    used_ids: set[str] = set()
    for row in payload.get("rewrites", []):
        if not isinstance(row, dict):
            continue
        evidence_id = str(row.get("evidence_id", ""))
        source = eligible.get(evidence_id)
        improved = clean_line(str(row.get("improved", "")))
        if not source or evidence_id in used_ids or not improved:
            continue
        if _contains_forbidden(improved, forbidden) or not _rewrite_is_grounded(improved, source["original"]):
            continue
        allowed_ids = {item["requirement_id"] for item in requirement_map[evidence_id]}
        requirement_ids = tuple(
            str(value) for value in row.get("requirement_ids", []) if str(value) in allowed_ids
        )
        reason = clean_line(str(row.get("reason", ""))) or "Improves clarity while preserving the original evidence."
        rewrites.append(
            BulletRewrite(
                evidence_id=evidence_id,
                original=source["original"],
                improved=improved,
                reason=reason,
                requirement_ids=requirement_ids,
            )
        )
        used_ids.add(evidence_id)
    return OptimizedResume(
        headline=headline,
        professional_summary=summary,
        skills=tuple(selected_skills[:30]),
        rewrites=tuple(rewrites),
        mode="AI-tailored with evidence, numeric, and specificity validation",
    )


def optimize_resume_content(
    candidate: CandidateProfile,
    report: MatchReport,
    job_title: str,
    company: str,
    settings: Settings,
    use_ai: bool,
) -> OptimizedResume:
    if use_ai and settings.api_key:
        try:
            return ai_optimized_resume(candidate, report, job_title, company, settings)
        except Exception as exc:
            return deterministic_optimized_resume(
                candidate,
                report,
                job_title,
                warning=f"AI resume tailoring was unavailable; a conservative version was generated ({exc}).",
            )
    return deterministic_optimized_resume(candidate, report, job_title)


def optimized_resume_sections(
    candidate: CandidateProfile,
    optimized: OptimizedResume,
) -> tuple[list[str], list[tuple[str, list[str]]]]:
    original_sections = sectionize_resume(candidate.raw_text)
    header = list(next((lines for name, lines in original_sections if name == "Header"), []))
    has_recognized_sections = any(name != "Header" for name, _lines in original_sections)
    rewrite_lookup = {normalize(item.original): item.improved for item in optimized.rewrites}
    output: list[tuple[str, list[str]]] = []
    if optimized.professional_summary:
        output.append(("Professional Summary", [optimized.professional_summary]))
    if optimized.skills:
        output.append(("Core Skills & Tools", [" | ".join(optimized.skills)]))
    for section, lines in original_sections:
        if section in {"Header", "Professional Summary", "Skills", "Tools"}:
            continue
        revised_lines = [rewrite_lookup.get(normalize(clean_line(line)), clean_line(line)) for line in lines]
        revised_lines = [line for line in revised_lines if line]
        if revised_lines:
            output.append((section, revised_lines))
    if not has_recognized_sections:
        headline_key = normalize(_existing_headline(candidate))
        body = [
            clean_line(line)
            for line in header
            if clean_line(line) != candidate.candidate_name
            and not _is_contact_line(clean_line(line))
            and normalize(clean_line(line)) != headline_key
            and clean_line(line) != candidate.professional_summary
        ]
        body = [rewrite_lookup.get(normalize(line), line) for line in body if line]
        header = [
            line
            for line in header
            if clean_line(line) == candidate.candidate_name or _is_contact_line(clean_line(line))
        ]
        if body:
            output.append(("Experience & Qualifications", body))
    return header, output


def optimized_resume_markdown(candidate: CandidateProfile, optimized: OptimizedResume) -> str:
    header, sections = optimized_resume_sections(candidate, optimized)
    lines = [f"# {candidate.candidate_name}"]
    remaining_header = [clean_line(line) for line in header if clean_line(line) != candidate.candidate_name]
    if remaining_header:
        lines.append("  \n".join(remaining_header))
    if optimized.headline:
        lines.extend(["", f"**{optimized.headline}**"])
    for section, section_lines in sections:
        lines.extend(["", f"## {section}"])
        for line in section_lines:
            if section in {"Experience", "Projects"} and not _is_resume_context_line(line):
                lines.append(f"- {line}")
            else:
                lines.append(line)
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Safe exports
# ---------------------------------------------------------------------------


def report_csv(report: MatchReport) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(
        [
            "Requirement ID",
            "Requirement",
            "Category",
            "Importance",
            "Status",
            "Confidence",
            "Assessment",
            "Resume Evidence",
            "Evidence Source",
        ]
    )
    for item in report.matches:
        writer.writerow(
            [
                item.requirement.requirement_id,
                item.requirement.text,
                item.requirement.category.value,
                item.requirement.importance.value,
                item.status.value,
                f"{item.confidence:.2f}",
                item.assessment,
                item.evidence.text if item.evidence else "",
                item.evidence.source if item.evidence else "",
            ]
        )
    return buffer.getvalue().encode("utf-8-sig")


def report_pdf(candidate: CandidateProfile, job_title: str, company: str, report: MatchReport) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise RuntimeError("PDF export requires reportlab: pip install reportlab") from exc

    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=0.55 * inch,
        leftMargin=0.55 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.55 * inch,
        title="HireSense Explainable Match Report",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Score", parent=styles["Title"], fontSize=28, textColor=colors.HexColor("#1D4ED8"), alignment=TA_CENTER, spaceAfter=4))
    styles.add(ParagraphStyle(name="Small", parent=styles["BodyText"], fontSize=8, leading=10, textColor=colors.HexColor("#475569")))
    story: list[Any] = [
        Paragraph("EXPLAINABLE AI MATCH ASSESSMENT", styles["Small"]),
        Paragraph("HireSense Explainable Match Report", styles["Title"]),
        Paragraph(f"{report.score}%", styles["Score"]),
        Paragraph(report.alignment, styles["Heading2"]),
        Paragraph(
            html.escape(" · ".join(value for value in (job_title, company, candidate.candidate_name) if value)),
            styles["BodyText"],
        ),
        Spacer(1, 10),
    ]
    counts = Counter(item.status for item in report.matches)
    summary_data = [
        ["Direct Matches", "Related Evidence", "Uncertain", "Missing"],
        [
            str(counts[MatchStatus.DIRECT]),
            str(counts[MatchStatus.RELATED]),
            str(counts[MatchStatus.UNCERTAIN]),
            str(counts[MatchStatus.MISSING]),
        ],
    ]
    table = Table(summary_data, colWidths=[1.7 * inch] * 4)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EFF6FF")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFDBFE")),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.extend([Paragraph("Requirement Summary", styles["Heading2"]), table, Spacer(1, 10)])
    category_data: list[list[Any]] = [["Category", "Score", "Requirements"]]
    for item in report.category_scores:
        category_data.append(
            [item.category.value.replace("_", " ").title(), f"{item.percentage}%", str(item.requirements)]
        )
    category_table = Table(category_data, colWidths=[3.8 * inch, 1.2 * inch, 1.4 * inch])
    category_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DBEAFE")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.extend([Paragraph("Match by Category", styles["Heading2"]), category_table, PageBreak()])
    icons = {
        MatchStatus.DIRECT: "Direct Match",
        MatchStatus.RELATED: "Related Evidence",
        MatchStatus.UNCERTAIN: "Uncertain",
        MatchStatus.MISSING: "Missing",
    }
    for index, item in enumerate(report.matches, start=1):
        story.append(
            Paragraph(
                f"{index}. {html.escape(item.requirement.text)} — {icons[item.status]}",
                styles["Heading3"],
            )
        )
        story.append(
            Paragraph(
                f"Category: {item.requirement.category.value.replace('_', ' ').title()} &nbsp;&nbsp; "
                f"Importance: {item.requirement.importance.value.title()} &nbsp;&nbsp; "
                f"AI confidence: {item.confidence * 100:.0f}%",
                styles["Small"],
            )
        )
        story.append(Paragraph(f"<b>Assessment:</b> {html.escape(item.assessment)}", styles["BodyText"]))
        if item.evidence:
            story.append(
                Paragraph(
                    f"<b>Resume evidence:</b> {html.escape(item.evidence.text)}<br/>"
                    f"<font color='#64748B'>Evidence source: {html.escape(item.evidence.source)}</font>",
                    styles["BodyText"],
                )
            )
        else:
            story.append(Paragraph("<b>Resume evidence:</b> Not verified", styles["BodyText"]))
        story.append(Spacer(1, 9))
    document.build(story)
    return output.getvalue()


def safe_resume_stem(candidate: CandidateProfile, job_title: str) -> str:
    raw = f"{candidate.candidate_name}_{job_title}_Tailored_Resume".strip("_")
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", raw).strip("._-")
    return stem[:100] or "Tailored_Resume"


def improved_resume_docx(candidate: CandidateProfile, optimized: OptimizedResume) -> bytes:
    try:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError("Word export requires python-docx: pip install python-docx") from exc

    navy = RGBColor(15, 35, 68)
    blue = RGBColor(37, 99, 235)
    ink = RGBColor(31, 41, 55)
    muted = RGBColor(71, 85, 105)
    rule_color = "CBD5E1"

    def set_run_font(run: Any, size: float, color: Any = ink, bold: bool = False) -> None:
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.bold = bold
        fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attribute}"), "Arial")

    def paragraph_style(
        name: str,
        *,
        size: float,
        color: Any = ink,
        bold: bool = False,
        before: float = 0,
        after: float = 0,
        line_spacing: float = 1.08,
        keep_with_next: bool = False,
    ) -> Any:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = bold
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attribute}"), "Arial")
        fmt = style.paragraph_format
        fmt.space_before = Pt(before)
        fmt.space_after = Pt(after)
        fmt.line_spacing = line_spacing
        fmt.keep_with_next = keep_with_next
        fmt.widow_control = True
        return style

    def add_bottom_rule(style: Any) -> None:
        p_pr = style._element.get_or_add_pPr()
        borders = p_pr.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            p_pr.append(borders)
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "3")
        bottom.set(qn("w:color"), rule_color)
        borders.append(bottom)

    def create_bullet_numbering() -> int:
        numbering = document.part.numbering_part.element
        abstract_ids = [
            int(node.get(qn("w:abstractNumId")))
            for node in numbering.findall(qn("w:abstractNum"))
        ]
        num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
        abstract_id = max(abstract_ids, default=-1) + 1
        num_id = max(num_ids, default=0) + 1

        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        for tag, value in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", "•"), ("w:lvlJc", "left")):
            node = OxmlElement(tag)
            node.set(qn("w:val"), value)
            level.append(node)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), "540")
        indent.set(qn("w:hanging"), "270")
        p_pr.extend([tabs, indent])
        level.append(p_pr)
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), "Arial")
        r_fonts.set(qn("w:hAnsi"), "Arial")
        r_pr.append(r_fonts)
        level.append(r_pr)
        abstract.append(level)
        numbering.append(abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)
        return num_id

    def apply_bullet(paragraph: Any, num_id: int) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        level = OxmlElement("w:ilvl")
        level.set(qn("w:val"), "0")
        number = OxmlElement("w:numId")
        number.set(qn("w:val"), str(num_id))
        num_pr.extend([level, number])
        p_pr.append(num_pr)

    def add_page_field(paragraph: Any) -> None:
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        run = OxmlElement("w:r")
        run_properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Arial")
        fonts.set(qn("w:hAnsi"), "Arial")
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "16")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "64748B")
        run_properties.extend([fonts, size, color])
        text = OxmlElement("w:t")
        text.text = "1"
        run.extend([run_properties, text])
        field.append(run)
        paragraph._p.append(field)

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    # Resume-specific compact override to the compact_reference_guide preset.
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.7)
    normal.font.color.rgb = ink
    normal_fonts = normal._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        normal_fonts.set(qn(f"w:{attribute}"), "Arial")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True

    section_style = paragraph_style(
        "Resume Section", size=10.5, color=navy, bold=True, before=8, after=4, keep_with_next=True
    )
    add_bottom_rule(section_style)
    paragraph_style("Resume Body", size=9.7, after=2.5, line_spacing=1.08)
    paragraph_style(
        "Resume Role", size=9.8, color=navy, bold=True, before=4.5, after=1.5, keep_with_next=True
    )
    bullet_style = paragraph_style("Resume Bullet", size=9.45, after=1.8, line_spacing=1.07)
    bullet_style.paragraph_format.left_indent = Inches(0.375)
    bullet_style.paragraph_format.first_line_indent = Inches(-0.188)
    bullet_style.paragraph_format.keep_together = True
    paragraph_style("Resume Skills", size=9.35, after=2, line_spacing=1.08)
    bullet_num_id = create_bullet_numbering()

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(1.5)
    name.paragraph_format.keep_with_next = True
    run = name.add_run(candidate.candidate_name)
    set_run_font(run, 20, navy, True)

    header, sections = optimized_resume_sections(candidate, optimized)
    remaining_header = [clean_line(line) for line in header if clean_line(line) != candidate.candidate_name]
    if remaining_header:
        contact = document.add_paragraph(" | ".join(remaining_header))
        contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact.paragraph_format.space_after = Pt(2.5)
        contact.paragraph_format.line_spacing = 1.0
        contact.paragraph_format.keep_with_next = True
        for contact_run in contact.runs:
            set_run_font(contact_run, 8.4, muted)
    if optimized.headline:
        headline = document.add_paragraph(optimized.headline)
        headline.alignment = WD_ALIGN_PARAGRAPH.CENTER
        headline.paragraph_format.space_after = Pt(7)
        headline.paragraph_format.keep_with_next = True
        set_run_font(headline.runs[0], 10.4, blue, True)

    for section_name, lines in sections:
        heading = document.add_paragraph(style="Resume Section")
        heading.add_run(section_name.upper())
        for line in lines:
            if section_name in {"Experience", "Projects"} and not _is_resume_context_line(line):
                paragraph = document.add_paragraph(style="Resume Bullet")
                apply_bullet(paragraph, bullet_num_id)
                paragraph.add_run(line)
            else:
                if section_name in {"Experience", "Projects"} and _is_resume_context_line(line):
                    paragraph = document.add_paragraph(line, style="Resume Role")
                elif section_name == "Core Skills & Tools":
                    paragraph = document.add_paragraph(line, style="Resume Skills")
                else:
                    paragraph = document.add_paragraph(line, style="Resume Body")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    footer.paragraph_format.line_spacing = 1.0
    footer_run = footer.add_run(f"{candidate.candidate_name}  |  Page ")
    set_run_font(footer_run, 8, muted)
    add_page_field(footer)

    document.core_properties.title = f"{candidate.candidate_name} - Tailored Resume"
    document.core_properties.subject = "Evidence-grounded tailored resume generated by HireSense"
    document.core_properties.author = candidate.candidate_name
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def improved_resume_pdf(candidate: CandidateProfile, optimized: OptimizedResume) -> bytes:
    try:
        import reportlab
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise RuntimeError("PDF export requires reportlab: pip install reportlab") from exc

    # Embed ReportLab's bundled TrueType font so spacing and alignment remain
    # identical in browsers, Preview, Acrobat, and applicant-tracking systems.
    font_directory = os.path.join(os.path.dirname(reportlab.__file__), "fonts")
    regular_font = "HireSenseSans"
    bold_font = "HireSenseSans-Bold"
    pdfmetrics.registerFont(TTFont(regular_font, os.path.join(font_directory, "Vera.ttf")))
    pdfmetrics.registerFont(TTFont(bold_font, os.path.join(font_directory, "VeraBd.ttf")))

    navy = colors.HexColor("#0F2344")
    blue = colors.HexColor("#2563EB")
    ink = colors.HexColor("#1F2937")
    muted = colors.HexColor("#475569")
    rule = colors.HexColor("#CBD5E1")
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=0.55 * inch,
        leftMargin=0.55 * inch,
        topMargin=0.45 * inch,
        bottomMargin=0.50 * inch,
        title=f"{candidate.candidate_name} - Tailored Resume",
        author=candidate.candidate_name,
        subject="Evidence-grounded tailored resume generated by HireSense",
        allowSplitting=1,
    )
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="ResumeName",
            parent=styles["Title"],
            fontName=bold_font,
            fontSize=20,
            leading=22,
            alignment=TA_CENTER,
            textColor=navy,
            spaceAfter=1.5,
            keepWithNext=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ResumeContact",
            parent=styles["BodyText"],
            fontName=regular_font,
            fontSize=7.9,
            leading=9.2,
            alignment=TA_CENTER,
            textColor=muted,
            spaceAfter=2.5,
            keepWithNext=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ResumeHeadline",
            parent=styles["BodyText"],
            fontName=bold_font,
            fontSize=9.7,
            leading=11.2,
            alignment=TA_CENTER,
            textColor=blue,
            spaceAfter=7,
            keepWithNext=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ResumeSection",
            parent=styles["Heading2"],
            fontName=bold_font,
            fontSize=10,
            leading=11.2,
            textColor=navy,
            spaceBefore=5.5,
            spaceAfter=1.5,
            keepWithNext=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ResumeBody",
            parent=styles["BodyText"],
            fontName=regular_font,
            fontSize=8.7,
            leading=9.9,
            textColor=ink,
            alignment=TA_LEFT,
            spaceAfter=1.8,
            splitLongWords=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ResumeRole",
            parent=styles["BodyText"],
            fontName=bold_font,
            fontSize=8.85,
            leading=10.1,
            textColor=navy,
            spaceBefore=3.2,
            spaceAfter=1,
            keepWithNext=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ResumeBullet",
            parent=styles["BodyText"],
            fontName=regular_font,
            fontSize=8.55,
            leading=9.8,
            textColor=ink,
            leftIndent=27,
            firstLineIndent=0,
            bulletIndent=13.5,
            bulletFontName=regular_font,
            bulletFontSize=7.8,
            spaceAfter=1.1,
            splitLongWords=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ResumeSkills",
            parent=styles["BodyText"],
            fontName=regular_font,
            fontSize=8.6,
            leading=9.9,
            textColor=ink,
            spaceAfter=1.5,
            splitLongWords=True,
        )
    )

    def line_flowable(section_name: str, line: str) -> Any:
        escaped = html.escape(line)
        if section_name in {"Experience", "Projects"} and not _is_resume_context_line(line):
            return Paragraph(escaped, styles["ResumeBullet"], bulletText="•")
        if section_name in {"Experience", "Projects"}:
            return Paragraph(escaped, styles["ResumeRole"])
        if section_name == "Core Skills & Tools":
            return Paragraph(escaped, styles["ResumeSkills"])
        return Paragraph(escaped, styles["ResumeBody"])

    def draw_footer(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setStrokeColor(rule)
        canvas.setLineWidth(0.45)
        canvas.line(doc.leftMargin, 0.36 * inch, letter[0] - doc.rightMargin, 0.36 * inch)
        canvas.setFillColor(muted)
        canvas.setFont(regular_font, 7.8)
        canvas.drawRightString(
            letter[0] - doc.rightMargin,
            0.22 * inch,
            f"{candidate.candidate_name}  |  Page {canvas.getPageNumber()}",
        )
        canvas.restoreState()

    story: list[Any] = [Paragraph(html.escape(candidate.candidate_name), styles["ResumeName"])]
    header, sections = optimized_resume_sections(candidate, optimized)
    remaining_header = [clean_line(line) for line in header if clean_line(line) != candidate.candidate_name]
    if remaining_header:
        story.append(Paragraph(html.escape(" | ".join(remaining_header)), styles["ResumeContact"]))
    if optimized.headline:
        story.append(Paragraph(html.escape(optimized.headline), styles["ResumeHeadline"]))
    for section_name, lines in sections:
        section_header: list[Any] = [
            Paragraph(html.escape(section_name.upper()), styles["ResumeSection"]),
            HRFlowable(width="100%", thickness=0.6, color=rule, spaceBefore=0, spaceAfter=3),
        ]
        flowable_groups: list[list[Any]] = []
        position = 0
        while position < len(lines):
            line = lines[position]
            if (
                section_name in {"Experience", "Projects"}
                and _is_resume_context_line(line)
                and position + 1 < len(lines)
                and not _is_resume_context_line(lines[position + 1])
            ):
                flowable_groups.append(
                    [line_flowable(section_name, line), line_flowable(section_name, lines[position + 1])]
                )
                position += 2
            else:
                flowable_groups.append([line_flowable(section_name, line)])
                position += 1
        if flowable_groups:
            # Flatten the first group into the section wrapper. Nested
            # KeepTogether objects can incorrectly reserve a full page.
            story.append(KeepTogether(section_header + flowable_groups[0]))
            for group in flowable_groups[1:]:
                story.append(KeepTogether(group) if len(group) > 1 else group[0])
        else:
            story.extend(section_header)
    document.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
    return output.getvalue()


# ---------------------------------------------------------------------------
# Streamlit interface
# ---------------------------------------------------------------------------


def inject_styles() -> None:
    st.markdown(
        """
        <style>
        :root {
            --navy:#0b1f3a; --ink:#14213d; --blue:#2563eb; --cyan:#06b6d4; --green:#059669;
            --soft:#f5f8ff; --line:#dce5f3; --muted:#64748b; --white:#ffffff;
            --shadow:0 18px 45px rgba(15,35,68,.09);
        }
        @keyframes hs-enter { from { opacity:0; transform:translateY(18px); } to { opacity:1; transform:translateY(0); } }
        @keyframes hs-pop { from { opacity:0; transform:scale(.96); } to { opacity:1; transform:scale(1); } }
        @keyframes hs-float { 0%,100% { transform:translate3d(0,0,0); } 50% { transform:translate3d(-10px,12px,0); } }
        @keyframes hs-glow { 0%,100% { box-shadow:0 0 0 4px rgba(52,211,153,.14); } 50% { box-shadow:0 0 0 8px rgba(52,211,153,.05); } }
        @keyframes hs-bar { from { transform:scaleX(0); } to { transform:scaleX(1); } }
        @keyframes hs-gradient { 0% { background-position:0% 50%; } 100% { background-position:100% 50%; } }
        html,body,[class*="css"] { font-family:Inter,ui-sans-serif,-apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif; }
        html { scroll-behavior:smooth; }
        .stApp {
            background:radial-gradient(circle at 88% 3%,rgba(6,182,212,.12),transparent 24rem),
                       radial-gradient(circle at 8% 18%,rgba(37,99,235,.10),transparent 28rem),
                       linear-gradient(180deg,#f7faff 0,#ffffff 34rem);
        }
        [data-testid="stHeader"] { background:rgba(247,250,255,.72); backdrop-filter:blur(14px); }
        .block-container { max-width:1240px; padding-top:1.1rem; padding-bottom:4rem; }
        h1,h2,h3 { color:var(--navy); letter-spacing:-.035em; }
        h2 { margin-top:.2rem; }
        p { line-height:1.62; }
        [data-testid="stSidebar"] {
            background:linear-gradient(180deg,#0b1f3a 0%,#102d52 58%,#0c3b55 100%);
            border-right:1px solid rgba(255,255,255,.08);
        }
        [data-testid="stSidebar"] h1,[data-testid="stSidebar"] h2,[data-testid="stSidebar"] h3,
        [data-testid="stSidebar"] label,[data-testid="stSidebar"] p { color:#f8fbff; }
        [data-testid="stSidebar"] hr { border-color:rgba(255,255,255,.14); }
        [data-testid="stSidebar"] [data-testid="stAlert"] p { color:inherit; }
        .hs-side-brand { display:flex; align-items:center; gap:.72rem; margin:.15rem 0 1.15rem; }
        .hs-side-mark { width:2.35rem; height:2.35rem; border-radius:.82rem; display:grid; place-items:center; color:white; font-size:1.2rem; font-weight:900; background:linear-gradient(135deg,#3b82f6,#06b6d4); box-shadow:0 10px 24px rgba(6,182,212,.28); animation:hs-pop .55s ease-out both; }
        .hs-side-name { color:#fff; font-size:1.05rem; font-weight:800; line-height:1.05; }
        .hs-side-caption { color:#a8c4e5; font-size:.73rem; margin-top:.18rem; }
        .hs-side-guide { border:1px solid rgba(255,255,255,.12); border-radius:16px; padding:1rem; background:rgba(255,255,255,.06); }
        .hs-side-guide strong { color:#fff; display:block; font-size:.92rem; margin-bottom:.65rem; }
        .hs-side-step { display:flex; gap:.62rem; align-items:flex-start; color:#d7e8f8; font-size:.79rem; line-height:1.42; margin:.58rem 0; }
        .hs-side-step span { display:grid; place-items:center; flex:0 0 auto; width:1.45rem; height:1.45rem; border-radius:50%; color:#fff; font-size:.7rem; font-weight:850; background:linear-gradient(135deg,#2563eb,#06b6d4); }
        .hs-side-trust { color:#a8c4e5; font-size:.73rem; line-height:1.5; margin-top:1rem; }
        .hs-hero { position:relative; overflow:hidden; color:#fff; border-radius:28px; padding:2.15rem 2.3rem 2rem; background:linear-gradient(125deg,#081d36 0%,#123e70 55%,#066783 100%); box-shadow:0 28px 70px rgba(15,35,68,.20); margin:.35rem 0 1.35rem; animation:hs-pop .58s cubic-bezier(.2,.75,.25,1) both; }
        .hs-hero:before { content:""; position:absolute; inset:0; opacity:.12; background-image:radial-gradient(rgba(255,255,255,.8) .7px,transparent .7px); background-size:20px 20px; mask-image:linear-gradient(110deg,transparent 10%,#000 90%); }
        .hs-hero:after { content:""; position:absolute; width:25rem; height:25rem; border-radius:50%; right:-8rem; top:-13rem; background:radial-gradient(circle,rgba(103,232,249,.28),rgba(59,130,246,.04) 62%,transparent 70%); animation:hs-float 9s ease-in-out infinite; }
        .hs-hero-top { display:flex; align-items:center; justify-content:space-between; gap:1rem; position:relative; z-index:1; }
        .hs-brand-chip,.hs-privacy-chip { display:inline-flex; align-items:center; gap:.42rem; border:1px solid rgba(255,255,255,.18); background:rgba(255,255,255,.09); border-radius:999px; padding:.42rem .72rem; color:#dff8ff; font-size:.76rem; font-weight:750; letter-spacing:.035em; }
        .hs-live-dot { width:.48rem; height:.48rem; background:#34d399; border-radius:50%; box-shadow:0 0 0 4px rgba(52,211,153,.14); animation:hs-glow 2.8s ease-in-out infinite; }
        .hs-hero h1 { color:#fff; max-width:900px; font-size:clamp(2.1rem,4.4vw,3.55rem); line-height:1.02; margin:1.25rem 0 .75rem; position:relative; z-index:1; animation:hs-enter .62s .08s ease-out both; }
        .hs-hero h1 span { color:#67e8f9; text-shadow:0 0 28px rgba(103,232,249,.24); }
        .hs-hero-copy { max-width:760px; color:#d8e8fb; font-size:1.06rem; margin:0; position:relative; z-index:1; animation:hs-enter .62s .16s ease-out both; }
        .hs-hero-pills { display:flex; flex-wrap:wrap; gap:.58rem; margin-top:1.35rem; position:relative; z-index:1; animation:hs-enter .62s .24s ease-out both; }
        .hs-hero-pills span { background:rgba(255,255,255,.10); border:1px solid rgba(255,255,255,.13); border-radius:10px; padding:.48rem .7rem; color:#edf7ff; font-size:.78rem; font-weight:650; }
        .hs-flow { display:grid; grid-template-columns:repeat(3,1fr); gap:.8rem; margin:.85rem 0 1.5rem; }
        .hs-flow-card { background:rgba(255,255,255,.88); border:1px solid var(--line); border-radius:16px; padding:1rem 1.05rem; box-shadow:0 8px 24px rgba(15,35,68,.045); transition:transform .22s ease,border-color .22s ease,box-shadow .22s ease; animation:hs-enter .55s ease-out both; }
        .hs-flow-card:nth-child(2) { animation-delay:.09s; }
        .hs-flow-card:nth-child(3) { animation-delay:.18s; }
        .hs-flow-card:hover { transform:translateY(-4px); border-color:#b8d6ee; box-shadow:0 16px 34px rgba(15,35,68,.09); }
        .hs-flow-number { color:var(--blue); font-size:.72rem; font-weight:900; letter-spacing:.12em; }
        .hs-flow-card strong { color:var(--navy); display:block; margin:.22rem 0 .12rem; }
        .hs-flow-card small { color:var(--muted); line-height:1.45; }
        .hs-value-grid { display:grid; grid-template-columns:repeat(3,1fr); gap:.8rem; margin:.4rem 0 1.35rem; }
        .hs-value-card { border:1px solid var(--line); border-radius:17px; padding:1rem 1.05rem; background:linear-gradient(145deg,#fff,#f8fbff); box-shadow:0 8px 22px rgba(15,35,68,.045); }
        .hs-value-icon { width:2.05rem; height:2.05rem; display:grid; place-items:center; border-radius:.7rem; color:#fff; font-weight:900; background:linear-gradient(135deg,#2563eb,#0d9488); margin-bottom:.62rem; }
        .hs-value-card strong { display:block; color:var(--navy); font-size:.95rem; margin-bottom:.24rem; }
        .hs-value-card span { display:block; color:var(--muted); font-size:.8rem; line-height:1.5; }
        .stTabs [data-baseweb="tab-list"] { gap:.42rem; padding:.42rem; border:1px solid #cbdcf2; border-radius:16px; background:linear-gradient(120deg,#edf5ff,#eefbf7); box-shadow:0 10px 28px rgba(15,35,68,.07); flex-wrap:wrap; }
        .stTabs [data-baseweb="tab"] { min-height:3.05rem; border:1px solid transparent; border-radius:12px; padding:0 1.05rem; color:#334155; font-weight:780; cursor:pointer; transition:background .2s ease,color .2s ease,transform .2s ease,box-shadow .2s ease; }
        .stTabs [data-baseweb="tab"]:hover { color:#1d4ed8; background:rgba(255,255,255,.75); transform:translateY(-1px); }
        .stTabs [aria-selected="true"] { border-color:transparent; background:linear-gradient(105deg,#1d4ed8,#0284c7)!important; color:#fff!important; box-shadow:0 8px 20px rgba(37,99,235,.25); }
        .stTabs [data-baseweb="tab-highlight"] { display:none; }
        [data-testid="stMetric"] {
            background:linear-gradient(145deg,#fff,#f8fbff); border:1px solid var(--line); border-radius:16px;
            padding:15px 16px; box-shadow:0 8px 22px rgba(15,35,68,.05);
        }
        [data-testid="stMetricValue"] { color:var(--navy); font-weight:820; }
        [data-testid="stForm"] {
            background:rgba(255,255,255,.92); border:1px solid var(--line); border-radius:20px;
            padding:22px; box-shadow:var(--shadow);
        }
        [data-testid="stFileUploaderDropzone"] { background:#f8fbff; border:1.5px dashed #b7c8e5; border-radius:15px; }
        [data-testid="stExpander"] { background:rgba(255,255,255,.86); border:1px solid var(--line); border-radius:14px; overflow:hidden; box-shadow:0 5px 16px rgba(15,35,68,.035); }
        div[data-testid="stAlert"] { border-radius:13px; border-width:1px; }
        [data-testid="stBaseButton-primary"],[data-testid="stBaseButton-primaryFormSubmit"] {
            border:0!important; color:#fff!important; font-weight:780!important;
            background:linear-gradient(105deg,#1d4ed8,#0284c7,#0e7490)!important; background-size:180% 180%!important;
            box-shadow:0 10px 24px rgba(37,99,235,.24)!important;
            transition:transform .2s ease,box-shadow .2s ease,filter .2s ease!important;
            animation:hs-gradient 3.8s ease-in-out infinite alternate;
        }
        [data-testid="stBaseButton-secondary"],[data-testid="stBaseButton-secondaryFormSubmit"],.stDownloadButton button {
            border:0!important; color:#fff!important; font-weight:760!important;
            background:linear-gradient(105deg,#047857,#059669,#0d9488)!important; background-size:180% 180%!important;
            box-shadow:0 9px 22px rgba(5,150,105,.20)!important;
            transition:transform .2s ease,box-shadow .2s ease,filter .2s ease!important;
        }
        .stLinkButton a { border:0!important; color:#fff!important; background:linear-gradient(105deg,#1d4ed8,#0284c7)!important; box-shadow:0 9px 22px rgba(37,99,235,.18)!important; }
        [data-testid="stBaseButton-primary"]:hover,[data-testid="stBaseButton-primaryFormSubmit"]:hover,
        [data-testid="stBaseButton-secondary"]:hover,[data-testid="stBaseButton-secondaryFormSubmit"]:hover,
        .stDownloadButton button:hover,.stLinkButton a:hover { transform:translateY(-2px); filter:saturate(1.08) brightness(1.03); }
        [data-testid="stBaseButton-primary"]:hover,[data-testid="stBaseButton-primaryFormSubmit"]:hover { box-shadow:0 14px 30px rgba(37,99,235,.30)!important; }
        [data-testid="stBaseButton-secondary"]:hover,[data-testid="stBaseButton-secondaryFormSubmit"]:hover,.stDownloadButton button:hover { box-shadow:0 13px 28px rgba(5,150,105,.28)!important; }
        [data-testid^="stBaseButton-"]:focus-visible,.stLinkButton a:focus-visible { outline:3px solid rgba(14,165,233,.32)!important; outline-offset:2px; }
        .stDownloadButton button,.stLinkButton a,.stButton button,.stFormSubmitButton button { border-radius:11px!important; font-weight:700; }
        .hs-eyebrow { color:#0e7490; font-weight:850; font-size:.74rem; letter-spacing:.13em; text-transform:uppercase; margin-top:.35rem; }
        .hs-subtle { color:var(--muted); margin-top:-.45rem; }
        .hs-score-panel { display:flex; align-items:center; gap:1.75rem; background:linear-gradient(135deg,#ffffff,#f1f7ff); border:1px solid #cedcf1; border-radius:23px; padding:1.55rem 1.75rem; box-shadow:var(--shadow); margin:.55rem 0 1.25rem; }
        .hs-score-ring { --score:0; width:8.25rem; height:8.25rem; border-radius:50%; flex:0 0 auto; display:grid; place-items:center; background:conic-gradient(#2563eb calc(var(--score)*.7%),#0d9488 calc(var(--score)*1%),#dce7f5 0); box-shadow:0 12px 30px rgba(37,99,235,.16); animation:hs-pop .6s ease-out both; }
        .hs-score-ring>div { width:6.45rem; height:6.45rem; display:grid; place-content:center; text-align:center; border-radius:50%; background:#fff; }
        .hs-score-ring strong { color:var(--navy); font-size:2.35rem; line-height:1; }
        .hs-score-ring span { color:var(--muted); font-size:.72rem; text-transform:uppercase; letter-spacing:.12em; margin-top:.3rem; }
        .hs-score-copy span { color:#0e7490; font-size:.72rem; font-weight:850; letter-spacing:.12em; text-transform:uppercase; }
        .hs-score-copy h3 { color:var(--navy); font-size:1.55rem; margin:.28rem 0 .32rem; }
        .hs-score-copy p { color:var(--muted); margin:0; max-width:570px; }
        .hs-category {
            background:linear-gradient(145deg,#fff,#f9fbff); border:1px solid var(--line); border-radius:15px;
            padding:15px 16px; min-height:126px; margin-bottom:12px; box-shadow:0 7px 20px rgba(15,35,68,.045);
            transition:transform .18s ease,box-shadow .18s ease;
        }
        .hs-category:hover { transform:translateY(-2px); box-shadow:0 11px 25px rgba(15,35,68,.08); }
        .hs-category-name { color:#334155; font-weight:760; font-size:.87rem; }
        .hs-category-score { color:var(--navy); font-weight:880; font-size:1.7rem; margin-top:.2rem; }
        .hs-category-count { color:var(--muted); font-size:.8rem; }
        .hs-category-bar { height:.38rem; margin-top:.62rem; border-radius:99px; background:#e5edf8; overflow:hidden; }
        .hs-category-bar span { display:block; height:100%; border-radius:99px; background:linear-gradient(90deg,#2563eb,#06b6d4,#10b981); transform-origin:left; animation:hs-bar .8s ease-out both; }
        .hs-report-kicker { display:flex; align-items:center; gap:.55rem; color:#0e7490; font-size:.74rem; font-weight:850; letter-spacing:.12em; text-transform:uppercase; margin:.15rem 0 .2rem; }
        .hs-report-kicker:before { content:""; width:1.8rem; height:3px; border-radius:99px; background:linear-gradient(90deg,#2563eb,#10b981); }
        .hs-focus-card { height:100%; min-height:148px; padding:1rem 1.05rem; border:1px solid var(--line); border-radius:16px; background:linear-gradient(145deg,#fff,#f8fbff); box-shadow:0 8px 22px rgba(15,35,68,.05); transition:transform .2s ease,box-shadow .2s ease; }
        .hs-focus-card:hover { transform:translateY(-3px); box-shadow:0 14px 30px rgba(15,35,68,.09); }
        .hs-focus-card.clarify { border-top:3px solid #2563eb; }
        .hs-focus-card.verify { border-top:3px solid #f59e0b; }
        .hs-focus-index { color:#64748b; font-size:.68rem; font-weight:850; letter-spacing:.12em; text-transform:uppercase; }
        .hs-focus-card strong { display:block; color:var(--navy); font-size:.95rem; line-height:1.35; margin:.38rem 0 .62rem; }
        .hs-focus-footer { display:flex; align-items:center; justify-content:space-between; gap:.5rem; color:#64748b; font-size:.73rem; }
        .hs-focus-type { border-radius:999px; padding:.25rem .5rem; font-weight:760; background:#eff6ff; color:#1d4ed8; }
        .hs-focus-card.verify .hs-focus-type { background:#fff7ed; color:#b45309; }
        .hs-compact-banner { display:flex; align-items:flex-start; justify-content:space-between; gap:1rem; border:1px solid #cfe0f4; border-radius:15px; background:linear-gradient(120deg,#f7fbff,#eefbf7); padding:.9rem 1rem; margin:.2rem 0 1rem; }
        .hs-compact-banner strong { color:var(--navy); }
        .hs-compact-banner span { color:#64748b; font-size:.83rem; line-height:1.45; }
        .hs-meta-row { display:flex; flex-wrap:wrap; gap:.42rem; margin:.25rem 0 .75rem; }
        .hs-meta-pill { display:inline-flex; align-items:center; border-radius:999px; padding:.3rem .58rem; font-size:.73rem; font-weight:720; color:#334155; background:#f1f5f9; border:1px solid #e2e8f0; }
        .hs-meta-pill.blue { color:#1d4ed8; background:#eff6ff; border-color:#bfdbfe; }
        .hs-meta-pill.green { color:#047857; background:#ecfdf5; border-color:#a7f3d0; }
        .hs-meta-pill.amber { color:#b45309; background:#fffbeb; border-color:#fde68a; }
        .hs-meta-pill.rose { color:#be123c; background:#fff1f2; border-color:#fecdd3; }
        .hs-evidence-quote { border-left:4px solid #2563eb; border-radius:0 12px 12px 0; background:#eff6ff; color:#174ea6; padding:.8rem .95rem; line-height:1.55; margin:.2rem 0 .45rem; }
        .hs-keywords { display:flex; flex-wrap:wrap; gap:.42rem; margin:.45rem 0 .8rem; }
        .hs-keywords span { border-radius:999px; padding:.32rem .62rem; color:#0f5e55; background:#ecfdf5; border:1px solid #a7f3d0; font-size:.75rem; font-weight:720; }
        .hs-results-note { color:#64748b; font-size:.82rem; margin:-.2rem 0 .7rem; }
        .hs-resume-banner { display:flex; justify-content:space-between; align-items:center; gap:1rem; border-radius:14px 14px 0 0; padding:.78rem 1rem; color:#ddecff; background:linear-gradient(120deg,#0b1f3a,#164e75); margin-top:.35rem; }
        .hs-resume-banner strong { color:#fff; }
        .hs-resume-banner span { font-size:.76rem; color:#bcd4ed; }
        .hs-footer { text-align:center; color:#7b8ba3; font-size:.77rem; padding:2rem 0 .5rem; }
        @media (max-width:760px) {
            .block-container { padding-left:1rem; padding-right:1rem; }
            .hs-hero { border-radius:21px; padding:1.55rem 1.25rem; }
            .hs-hero-top { align-items:flex-start; flex-direction:column; }
            .hs-flow { grid-template-columns:1fr; }
            .hs-value-grid { grid-template-columns:1fr; }
            .stTabs [data-baseweb="tab"] { flex:1 1 46%; padding:0 .7rem; }
            .hs-score-panel { align-items:flex-start; flex-direction:column; }
            .hs-score-ring { width:7.3rem; height:7.3rem; }
            .hs-score-ring>div { width:5.7rem; height:5.7rem; }
        }
        @media (prefers-reduced-motion:reduce) {
            html { scroll-behavior:auto; }
            *,*::before,*::after { animation-duration:.01ms!important; animation-iteration-count:1!important; transition-duration:.01ms!important; }
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


def initialize_state() -> None:
    defaults = {
        "candidate": None,
        "parsed_resume": None,
        "requirements": None,
        "report": None,
        "optimized_resume": None,
        "applications": [],
        "job_title": "",
        "company": "",
        "job_location": "",
        "job_description": "",
        "analysis_error": "",
    }
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value


def render_sidebar() -> tuple[Settings, bool]:
    settings = get_settings()
    with st.sidebar:
        st.markdown(
            """
            <div class="hs-side-brand">
              <div class="hs-side-mark">H</div>
              <div><div class="hs-side-name">HireSense</div><div class="hs-side-caption">Smarter, evidence-first applications</div></div>
            </div>
            <div class="hs-side-guide">
              <strong>Three steps to a stronger application</strong>
              <div class="hs-side-step"><span>1</span><div>Upload your résumé and paste a job description.</div></div>
              <div class="hs-side-step"><span>2</span><div>See your matches, evidence, and priority gaps.</div></div>
              <div class="hs-side-step"><span>3</span><div>Improve and download a tailored résumé.</div></div>
            </div>
            <div class="hs-side-trust">Your résumé is processed for the current session. HireSense never adds experience you did not provide.</div>
            """,
            unsafe_allow_html=True,
        )
    return settings, bool(settings.api_key)


def analyze_submission(
    upload: Any,
    job_description: str,
    job_title: str,
    company: str,
    job_location: str,
    settings: Settings,
    use_ai: bool,
) -> None:
    if len(job_description.strip()) < 100:
        raise ValueError("Paste at least 100 characters of the job description for a reliable comparison.")
    parsed = parse_resume(upload.name, upload.getvalue(), upload.type or "", settings)
    candidate = build_candidate_profile(parsed)
    requirements, extraction_mode, extraction_warnings = extract_requirements(job_description, settings, use_ai)
    report = create_match_report(
        candidate,
        requirements,
        settings,
        use_ai,
        initial_warnings=extraction_warnings,
        extraction_mode=extraction_mode,
    )
    st.session_state.parsed_resume = parsed
    st.session_state.candidate = candidate
    st.session_state.requirements = requirements
    st.session_state.report = report
    st.session_state.optimized_resume = None
    st.session_state.tailored_resume_consent = False
    st.session_state.job_title = job_title.strip()
    st.session_state.company = company.strip()
    st.session_state.job_location = job_location.strip()
    st.session_state.job_description = job_description
    st.session_state.analysis_error = ""


def render_analyze_input(settings: Settings, use_ai: bool) -> None:
    st.markdown('<div class="hs-eyebrow">Step 1 · Analyze</div>', unsafe_allow_html=True)
    st.header("Analyze your résumé against a job")
    st.caption("Upload your résumé and paste the complete job description. Every extracted requirement is evaluated independently.")
    with st.form("analysis_form", clear_on_submit=False, border=True):
        left, right = st.columns([0.9, 1.1], gap="large")
        with left:
            st.subheader("Résumé")
            upload = st.file_uploader(
                "Upload PDF or DOCX",
                type=["pdf", "docx"],
                help=f"Maximum {settings.max_upload_mb} MB. Scanned PDFs must contain OCR text.",
            )
            st.markdown(
                "The matcher preserves exact résumé excerpts and their section/role source. "
                "It never treats an AI-generated paraphrase as evidence."
            )
        with right:
            st.subheader("Job description")
            title_col, company_col, location_col = st.columns(3)
            job_title = title_col.text_input("Job title", value=st.session_state.job_title)
            company = company_col.text_input("Company", value=st.session_state.company)
            job_location = location_col.text_input("Location", value=st.session_state.job_location)
            job_description = st.text_area(
                "Complete job description",
                value=st.session_state.job_description,
                height=370,
                placeholder="Paste responsibilities, required qualifications, and preferred qualifications…",
            )
        submitted = st.form_submit_button("Create explainable match report", type="primary", width="stretch")

    if not submitted:
        return
    if upload is None:
        st.warning("Upload a PDF or DOCX résumé to continue.")
        return
    try:
        with st.status("Building the evidence map…", expanded=True) as status:
            st.write("Reading and structuring the résumé")
            st.write("Extracting atomic job requirements")
            st.write("Retrieving and validating evidence for each requirement")
            analyze_submission(
                upload,
                job_description,
                job_title,
                company,
                job_location,
                settings,
                use_ai,
            )
            status.update(label="Match report complete", state="complete", expanded=False)
    except (ResumeParserError, ValueError) as exc:
        st.session_state.analysis_error = str(exc)
        st.error(str(exc))
    except Exception as exc:
        st.session_state.analysis_error = f"{type(exc).__name__}: {exc}"
        st.error(f"The analysis could not be completed: {exc}")
        st.caption("Verify the API model/key, package installation, file readability, and job-description text.")


def _status_label(status: MatchStatus) -> str:
    return {
        MatchStatus.DIRECT: "✅ Direct Match",
        MatchStatus.RELATED: "🟡 Related Evidence",
        MatchStatus.UNCERTAIN: "⚠️ Uncertain",
        MatchStatus.MISSING: "❌ Missing",
    }[status]


def _category_label(category: Category) -> str:
    return category.value.replace("_", " ").title()


def _tracker_id(title: str, company: str, location: str) -> str:
    key = normalize("|".join((company, title, location))).replace(" ", "-")
    return key or f"application-{len(st.session_state.applications) + 1}"


def save_current_job_to_tracker(report: MatchReport, resume_version: str = "Original") -> tuple[bool, str]:
    title = st.session_state.job_title.strip() or "Pasted job description"
    company = st.session_state.company.strip()
    location = st.session_state.job_location.strip()
    record_id = _tracker_id(title, company, location)
    applications = list(st.session_state.applications)
    for record in applications:
        if record.get("id") == record_id:
            record.update(
                {
                    "match_score": report.score,
                    "required_score": report.required_score if report.required_score is not None else "",
                    "resume_version": resume_version,
                }
            )
            st.session_state.applications = applications
            return False, "This job was already tracked; its score and résumé version were updated."
    applications.append(
        {
            "id": record_id,
            "job_title": title,
            "company": company,
            "location": location,
            "status": "Saved",
            "application_date": "",
            "follow_up_date": "",
            "match_score": report.score,
            "required_score": report.required_score if report.required_score is not None else "",
            "resume_version": resume_version,
            "next_action": "Review tailored résumé and apply",
            "notes": "",
        }
    )
    st.session_state.applications = applications
    return True, "Job added to the application tracker."


def tracker_csv(applications: Sequence[dict[str, Any]]) -> bytes:
    columns = [
        "job_title",
        "company",
        "location",
        "status",
        "application_date",
        "follow_up_date",
        "match_score",
        "required_score",
        "resume_version",
        "next_action",
        "notes",
    ]
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    for row in applications:
        writer.writerow(row)
    return buffer.getvalue().encode("utf-8-sig")


def render_tracker() -> None:
    st.markdown('<div class="hs-eyebrow">Stage 2 · Track</div>', unsafe_allow_html=True)
    st.header("Application tracker")
    st.caption("Track the job, résumé version, follow-up date, next action, and application outcome.")

    report: MatchReport | None = st.session_state.report
    if report and st.button("Add the analyzed job to tracker", type="primary", width="stretch"):
        created, message = save_current_job_to_tracker(
            report,
            "Tailored" if st.session_state.optimized_resume else "Original",
        )
        (st.success if created else st.info)(message)

    with st.expander("Add an application manually", expanded=not st.session_state.applications):
        with st.form("manual_application_form", clear_on_submit=True):
            first = st.columns(3)
            title = first[0].text_input("Job title")
            company = first[1].text_input("Company")
            location = first[2].text_input("Location")
            second = st.columns(3)
            status = second[0].selectbox(
                "Status",
                ["Saved", "Applied", "Recruiter Screen", "Interviewing", "Offer", "Rejected", "Withdrawn", "Closed"],
            )
            application_date = second[1].date_input("Application date", value=date.today())
            follow_up_date = second[2].text_input("Follow-up date", placeholder="YYYY-MM-DD")
            notes = st.text_area("Notes", height=80)
            add_manual = st.form_submit_button("Add application", width="stretch")
        if add_manual:
            if not title.strip():
                st.warning("Enter a job title.")
            else:
                applications = list(st.session_state.applications)
                record_id = _tracker_id(title, company, location)
                if any(row.get("id") == record_id for row in applications):
                    st.warning("That job is already in the tracker.")
                else:
                    applications.append(
                        {
                            "id": record_id,
                            "job_title": title.strip(),
                            "company": company.strip(),
                            "location": location.strip(),
                            "status": status,
                            "application_date": application_date.isoformat(),
                            "follow_up_date": follow_up_date.strip(),
                            "match_score": "",
                            "required_score": "",
                            "resume_version": "Original",
                            "next_action": "",
                            "notes": notes.strip(),
                        }
                    )
                    st.session_state.applications = applications
                    st.success("Application added.")
                    st.rerun()

    applications = st.session_state.applications
    if not applications:
        st.info("No applications tracked yet. Add the analyzed job or create an entry manually.")
        return

    counts = Counter(str(row.get("status", "Saved")) for row in applications)
    metrics = st.columns(4)
    metrics[0].metric("Total", len(applications))
    metrics[1].metric("Applied", counts["Applied"])
    metrics[2].metric("Interviewing", counts["Recruiter Screen"] + counts["Interviewing"])
    metrics[3].metric("Offers", counts["Offer"])

    try:
        import pandas as pd
    except ImportError:
        for row in applications:
            st.write(
                f"**{row.get('job_title', 'Untitled')}** · {row.get('company', '')} · "
                f"{row.get('status', 'Saved')} · Score {row.get('match_score', '')}"
            )
        st.warning("Install pandas to edit the tracker table: pip install pandas")
    else:
        display_columns = [
            "id",
            "job_title",
            "company",
            "location",
            "status",
            "application_date",
            "follow_up_date",
            "match_score",
            "required_score",
            "resume_version",
            "next_action",
            "notes",
        ]
        frame = pd.DataFrame(applications)
        for column in display_columns:
            if column not in frame:
                frame[column] = ""
        for numeric_column in ("match_score", "required_score"):
            frame[numeric_column] = pd.to_numeric(frame[numeric_column], errors="coerce")
        edited = st.data_editor(
            frame[display_columns],
            width="stretch",
            hide_index=True,
            num_rows="dynamic",
            column_config={
                "id": None,
                "job_title": st.column_config.TextColumn("Job title", required=True),
                "company": st.column_config.TextColumn("Company"),
                "location": st.column_config.TextColumn("Location"),
                "status": st.column_config.SelectboxColumn(
                    "Status",
                    options=["Saved", "Applied", "Recruiter Screen", "Interviewing", "Offer", "Rejected", "Withdrawn", "Closed"],
                    required=True,
                ),
                "application_date": st.column_config.TextColumn("Applied", help="YYYY-MM-DD"),
                "follow_up_date": st.column_config.TextColumn("Follow-up", help="YYYY-MM-DD"),
                "match_score": st.column_config.NumberColumn("Match %", min_value=0, max_value=100),
                "required_score": st.column_config.NumberColumn("Required %", min_value=0, max_value=100),
                "resume_version": st.column_config.SelectboxColumn("Résumé", options=["Original", "Tailored"]),
                "next_action": st.column_config.TextColumn("Next action"),
                "notes": st.column_config.TextColumn("Notes"),
            },
            key="application_tracker_editor",
        )
        action_cols = st.columns(2)
        if action_cols[0].button("Save tracker changes", type="primary", width="stretch"):
            updated: list[dict[str, Any]] = []
            seen_ids: set[str] = set()
            for index, row in enumerate(edited.to_dict(orient="records"), start=1):
                cleaned: dict[str, Any] = {}
                for key, value in row.items():
                    if pd.isna(value):
                        value = ""
                    cleaned[key] = value.item() if hasattr(value, "item") else value
                if not str(cleaned.get("job_title", "")).strip():
                    continue
                cleaned["id"] = str(cleaned.get("id") or _tracker_id(
                    str(cleaned.get("job_title", "")),
                    str(cleaned.get("company", "")),
                    str(cleaned.get("location", "")),
                ))
                if cleaned["id"] in seen_ids:
                    continue
                seen_ids.add(cleaned["id"])
                updated.append(cleaned)
            st.session_state.applications = updated
            st.success("Tracker updated.")
            st.rerun()
        action_cols[1].download_button(
            "Download tracker · CSV",
            tracker_csv(edited.to_dict(orient="records")),
            file_name="hiresense_application_tracker.csv",
            mime="text/csv",
            width="stretch",
        )
    st.caption("Tracker entries are session-only. Download the CSV before closing the app if you want to keep them.")


def render_resume_preview(candidate: CandidateProfile, optimized: OptimizedResume) -> None:
    st.markdown(
        """
        <div class="hs-resume-banner">
          <strong>Tailored résumé preview</strong>
          <span>ATS-safe typography · aligned bullets · recruiter-ready hierarchy</span>
        </div>
        """,
        unsafe_allow_html=True,
    )
    with st.container(border=True):
        st.markdown(optimized_resume_markdown(candidate, optimized))


def _short_ui_text(value: str, limit: int = 112) -> str:
    cleaned = clean_line(value)
    if len(cleaned) <= limit:
        return cleaned
    shortened = cleaned[: limit - 1].rsplit(" ", 1)[0]
    return f"{shortened or cleaned[: limit - 1]}…"


def _compact_status(status: MatchStatus) -> tuple[str, str]:
    return {
        MatchStatus.DIRECT: ("Direct match", "green"),
        MatchStatus.RELATED: ("Related evidence", "blue"),
        MatchStatus.UNCERTAIN: ("Needs verification", "amber"),
        MatchStatus.MISSING: ("Missing evidence", "rose"),
    }[status]


def _render_report_overview(report: MatchReport) -> None:
    counts = Counter(item.status for item in report.matches)
    metrics = st.columns(4)
    metrics[0].metric("Direct matches", counts[MatchStatus.DIRECT])
    metrics[1].metric("Related evidence", counts[MatchStatus.RELATED])
    metrics[2].metric(
        "Needs attention",
        counts[MatchStatus.UNCERTAIN] + counts[MatchStatus.MISSING],
    )
    metrics[3].metric(
        "Required alignment",
        f"{report.required_score}%" if report.required_score is not None else "Not stated",
    )

    st.subheader("Match by category")
    category_columns = st.columns(3)
    for index, item in enumerate(report.category_scores):
        with category_columns[index % 3]:
            st.markdown(
                f"""
                <div class="hs-category">
                  <div class="hs-category-name">{html.escape(_category_label(item.category))}</div>
                  <div class="hs-category-score">{item.percentage}%</div>
                  <div class="hs-category-count">{item.requirements} requirement{'s' if item.requirements != 1 else ''}</div>
                  <div class="hs-category-bar"><span style="width:{item.percentage}%"></span></div>
                </div>
                """,
                unsafe_allow_html=True,
            )

    recommendations = build_improvement_recommendations(report)
    st.subheader("Next best actions")
    if not recommendations:
        st.success("Every extracted requirement already has direct résumé evidence.")
        return
    focus_items = recommendations[:3]
    focus_columns = st.columns(len(focus_items))
    for index, (column, suggestion) in enumerate(zip(focus_columns, focus_items), start=1):
        action_type = "Verify first" if suggestion.requires_verification else "Clarify evidence"
        tone = "verify" if suggestion.requires_verification else "clarify"
        with column:
            st.markdown(
                f"""
                <div class="hs-focus-card {tone}">
                  <div class="hs-focus-index">Priority {index} · {html.escape(suggestion.priority)}</div>
                  <strong>{html.escape(_short_ui_text(suggestion.requirement, 105))}</strong>
                  <div class="hs-focus-footer">
                    <span class="hs-focus-type">{action_type}</span>
                    <span>up to +{suggestion.potential_points:.1f}</span>
                  </div>
                </div>
                """,
                unsafe_allow_html=True,
            )
    st.caption("Open Priority gaps for the action plan, or Evidence explorer to inspect any requirement.")


def _render_priority_gaps(candidate: CandidateProfile, report: MatchReport) -> None:
    recommendations = build_improvement_recommendations(report)
    ceiling = related_clarity_ceiling(report)
    quick_wins = [item for item in recommendations if not item.requires_verification]
    verify_first = [item for item in recommendations if item.requires_verification]
    metrics = st.columns(3)
    metrics[0].metric("Current score", f"{report.score}%")
    metrics[1].metric(
        "Clarity opportunity",
        f"Up to {ceiling}%",
        delta=f"+{max(0, ceiling - report.score)}",
        help="A conservative what-if calculation using only related rows already backed by résumé evidence.",
    )
    metrics[2].metric("Verify before adding", len(verify_first))

    st.markdown(
        f"""
        <div class="hs-compact-banner">
          <div><strong>{len(quick_wins)} quick win{'s' if len(quick_wins) != 1 else ''}</strong><br>
          <span>Clarify evidence that already exists in the résumé.</span></div>
          <div><strong>{len(verify_first)} item{'s' if len(verify_first) != 1 else ''} to verify</strong><br>
          <span>Never add these merely as keywords.</span></div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    keywords = supported_keywords(candidate, report)
    if keywords:
        st.markdown("**Supported keywords to emphasize**")
        st.markdown(
            '<div class="hs-keywords">'
            + "".join(f"<span>{html.escape(keyword)}</span>" for keyword in keywords[:12])
            + "</div>",
            unsafe_allow_html=True,
        )

    if not recommendations:
        st.success("No improvement actions are needed for the extracted requirements.")
        return

    show_all = st.toggle(
        "Show every recommendation",
        value=False,
        key="show_all_improvement_recommendations",
    )
    visible = recommendations if show_all else recommendations[:5]
    st.markdown(
        f'<div class="hs-results-note">Showing {len(visible)} of {len(recommendations)} prioritized actions.</div>',
        unsafe_allow_html=True,
    )
    for index, suggestion in enumerate(visible, start=1):
        action_type = "Verify" if suggestion.requires_verification else "Clarify"
        title = (
            f"{index}. {suggestion.priority} · {action_type} · "
            f"{_short_ui_text(suggestion.requirement, 105)}"
        )
        with st.expander(title, expanded=False):
            tone = "amber" if suggestion.requires_verification else "blue"
            st.markdown(
                f"""
                <div class="hs-meta-row">
                  <span class="hs-meta-pill {tone}">{action_type}</span>
                  <span class="hs-meta-pill">{html.escape(suggestion.priority)} priority</span>
                  <span class="hs-meta-pill">Potential +{suggestion.potential_points:.1f}</span>
                  <span class="hs-meta-pill">{html.escape(suggestion.requirement_id)}</span>
                </div>
                """,
                unsafe_allow_html=True,
            )
            st.write(f"**Action:** {suggestion.action}")
            st.caption(suggestion.rationale)
            if suggestion.current_evidence:
                st.markdown(
                    f'<div class="hs-evidence-quote">{html.escape(suggestion.current_evidence)}</div>',
                    unsafe_allow_html=True,
                )
    if not show_all and len(recommendations) > len(visible):
        st.caption(f"Turn on Show every recommendation to review the remaining {len(recommendations) - len(visible)} items.")


def _render_evidence_explorer(report: MatchReport) -> None:
    st.write("Filter the evidence map, then open only the requirements you want to inspect.")
    controls = st.columns([1.15, 1, 1])
    view_mode = controls[0].radio(
        "Focus",
        ["Needs attention", "All requirements", "Direct matches"],
        horizontal=True,
        key="evidence_view_mode",
    )
    category_values = [item.category.value for item in report.category_scores]
    category_filter = controls[1].selectbox(
        "Category",
        ["All categories"] + category_values,
        format_func=lambda value: value if value == "All categories" else _category_label(Category(value)),
        key="evidence_category_filter",
    )
    importance_filter = controls[2].selectbox(
        "Importance",
        ["All", "Required", "Preferred"],
        key="evidence_importance_filter",
    )

    if view_mode == "Needs attention":
        matches = [item for item in report.matches if item.status != MatchStatus.DIRECT]
    elif view_mode == "Direct matches":
        matches = [item for item in report.matches if item.status == MatchStatus.DIRECT]
    else:
        matches = list(report.matches)
    if category_filter != "All categories":
        matches = [item for item in matches if item.requirement.category.value == category_filter]
    if importance_filter != "All":
        selected_importance = Importance(importance_filter.casefold())
        matches = [item for item in matches if item.requirement.importance == selected_importance]

    show_all = st.toggle("Show all matching requirements", value=False, key="show_all_evidence_rows")
    visible = matches if show_all else matches[:8]
    st.markdown(
        f'<div class="hs-results-note">Showing {len(visible)} of {len(matches)} matching requirements.</div>',
        unsafe_allow_html=True,
    )
    if not matches:
        st.info("No requirements match the selected filters.")
        return

    for index, item in enumerate(visible, start=1):
        status_label, tone = _compact_status(item.status)
        title = f"{status_label} · {_short_ui_text(item.requirement.text, 122)}"
        with st.expander(title, expanded=False):
            st.markdown(
                f"""
                <div class="hs-meta-row">
                  <span class="hs-meta-pill {tone}">{status_label}</span>
                  <span class="hs-meta-pill">{html.escape(_category_label(item.requirement.category))}</span>
                  <span class="hs-meta-pill">{html.escape(item.requirement.importance.value.title())}</span>
                  <span class="hs-meta-pill">{item.confidence * 100:.0f}% confidence</span>
                  <span class="hs-meta-pill">{html.escape(item.requirement.requirement_id)}</span>
                </div>
                """,
                unsafe_allow_html=True,
            )
            st.write(item.assessment)
            if item.evidence:
                st.markdown(
                    f'<div class="hs-evidence-quote">{html.escape(item.evidence.text)}</div>',
                    unsafe_allow_html=True,
                )
                st.caption(f"Evidence source: {item.evidence.source}")
            else:
                st.caption("No verified résumé excerpt was returned for this requirement.")
    if not show_all and len(matches) > len(visible):
        st.caption(f"Turn on Show all matching requirements to review the remaining {len(matches) - len(visible)} rows.")


def _render_tailor_and_exports(
    settings: Settings,
    use_ai: bool,
    candidate: CandidateProfile,
    report: MatchReport,
) -> None:
    st.subheader("Create a tailored résumé")
    st.write(
        "Improve recruiter readability and supported keyword placement without adding uncertain or missing experience."
    )
    consent = st.checkbox(
        "I consent to generate a tailored résumé and understand that I must review it before applying.",
        key="tailored_resume_consent",
    )
    if consent and st.button("Generate tailored résumé", type="primary", width="stretch"):
        with st.status("Tailoring the résumé from verified evidence…", expanded=True) as status:
            st.write("Prioritizing recruiter-relevant supported keywords")
            st.write("Improving the summary and evidence-backed bullets")
            st.write("Validating numbers, named tools, frameworks, and certifications")
            optimized = optimize_resume_content(
                candidate,
                report,
                st.session_state.job_title,
                st.session_state.company,
                settings,
                use_ai,
            )
            st.session_state.optimized_resume = optimized
            status.update(label="Tailored résumé ready for review", state="complete", expanded=False)

    optimized: OptimizedResume | None = st.session_state.optimized_resume
    if optimized:
        for warning in optimized.warnings:
            st.warning(warning)
        st.caption(f"Generation mode: {optimized.mode}")
        render_resume_preview(candidate, optimized)
        with st.expander("Review what changed", expanded=False):
            if optimized.professional_summary != candidate.professional_summary:
                st.write("**Professional summary was refined for the target role.**")
            if optimized.rewrites:
                for rewrite in optimized.rewrites:
                    st.write(f"**Original:** {rewrite.original}")
                    st.success(rewrite.improved)
                    st.caption(
                        f"{rewrite.reason}"
                        + (f" · Supports {', '.join(rewrite.requirement_ids)}" if rewrite.requirement_ids else "")
                    )
            else:
                st.info("No bullet was rewritten; the generated version uses conservative formatting and keyword ordering.")
        stem = safe_resume_stem(candidate, st.session_state.job_title)
        resume_downloads = st.columns(2)
        try:
            docx_bytes = improved_resume_docx(candidate, optimized)
        except Exception as exc:
            resume_downloads[0].warning(f"Word export unavailable: {exc}")
        else:
            resume_downloads[0].download_button(
                "Download tailored résumé · Word",
                docx_bytes,
                file_name=f"{stem}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                width="stretch",
            )
        try:
            resume_pdf = improved_resume_pdf(candidate, optimized)
        except Exception as exc:
            resume_downloads[1].warning(f"PDF export unavailable: {exc}")
        else:
            resume_downloads[1].download_button(
                "Download tailored résumé · PDF",
                resume_pdf,
                file_name=f"{stem}.pdf",
                mime="application/pdf",
                width="stretch",
            )
        if st.button("Save tailored application to tracker", key="save_tailored_to_tracker", width="stretch"):
            created, message = save_current_job_to_tracker(report, "Tailored")
            (st.success if created else st.info)(message)
        st.warning("Review the preview carefully. HireSense improves wording but cannot verify facts not stated in your résumé.")

    st.divider()
    st.subheader("Download the analysis")
    export_cols = st.columns(2)
    export_cols[0].download_button(
        "Download evidence map · CSV",
        report_csv(report),
        file_name="hiresense_requirement_evidence.csv",
        mime="text/csv",
        width="stretch",
    )
    try:
        pdf_bytes = report_pdf(candidate, st.session_state.job_title, st.session_state.company, report)
    except Exception as exc:
        export_cols[1].warning(f"PDF export unavailable: {exc}")
    else:
        export_cols[1].download_button(
            "Download match report · PDF",
            pdf_bytes,
            file_name="hiresense_match_report.pdf",
            mime="application/pdf",
            width="stretch",
        )


def render_report(settings: Settings, use_ai: bool) -> None:
    report: MatchReport | None = st.session_state.report
    candidate: CandidateProfile | None = st.session_state.candidate
    parsed: ParsedResume | None = st.session_state.parsed_resume
    if not report or not candidate or not parsed:
        return

    st.divider()
    st.markdown('<div class="hs-report-kicker">Explainable match dashboard</div>', unsafe_allow_html=True)
    st.header("HireSense Match Report")
    job_line = " · ".join(
        value
        for value in (
            st.session_state.job_title,
            st.session_state.company,
            st.session_state.job_location,
        )
        if value
    )
    if job_line:
        st.caption(job_line)
    st.markdown(
        f"""
        <div class="hs-score-panel">
          <div class="hs-score-ring" style="--score:{report.score}"><div><strong>{report.score}%</strong><span>match</span></div></div>
          <div class="hs-score-copy">
            <span>Overall résumé-to-job alignment</span>
            <h3>{html.escape(report.alignment)}</h3>
            <p>Built from independently assessed job requirements and the exact résumé evidence supporting each result.</p>
          </div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    st.markdown(
        f"""
        <div class="hs-compact-banner">
          <div><strong>{html.escape(parsed.filename)}</strong><br><span>{parsed.word_count:,} words · {len(report.matches)} requirements</span></div>
          <div><strong>{html.escape(report.mode)}</strong><br><span>Evidence-grounded assessment</span></div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    if report.warnings:
        with st.expander(f"Analysis notes · {len(report.warnings)}", expanded=False):
            for warning in report.warnings:
                st.warning(warning)

    overview_tab, gaps_tab, evidence_tab = st.tabs(
        ["Overview", "Priority gaps", "Evidence explorer"]
    )
    with overview_tab:
        if st.button("Save this job to application tracker", type="secondary", width="stretch"):
            created, message = save_current_job_to_tracker(
                report,
                "Tailored" if st.session_state.optimized_resume else "Original",
            )
            (st.success if created else st.info)(message)
        _render_report_overview(report)
    with gaps_tab:
        _render_priority_gaps(candidate, report)
    with evidence_tab:
        _render_evidence_explorer(report)


def render_tailor_workspace(settings: Settings, use_ai: bool) -> None:
    st.markdown('<div class="hs-eyebrow">Step 3 · Tailor & improve</div>', unsafe_allow_html=True)
    st.header("Tailor and improve your résumé")
    st.caption("Turn verified match evidence into a polished, recruiter-ready résumé without inventing experience.")
    report: MatchReport | None = st.session_state.report
    candidate: CandidateProfile | None = st.session_state.candidate
    if not report or not candidate:
        st.info("Complete an analysis first. Your tailored résumé tools will appear here after the match report is ready.")
        return
    _render_tailor_and_exports(settings, use_ai, candidate, report)


def render_discover() -> None:
    st.markdown('<div class="hs-eyebrow">Step 2 · Discover</div>', unsafe_allow_html=True)
    st.header("Discover targeted jobs")
    st.caption("Build direct search links, then paste a complete job description into Analyze.")
    with st.form("discovery_form", border=True):
        left, right = st.columns(2)
        titles = left.text_input("Target titles", placeholder="Business Process Analyst, Process Mining Consultant")
        locations = right.text_input("Locations", placeholder="Chicago, Remote, Bengaluru")
        submitted = st.form_submit_button("Build searches", type="primary", width="stretch")
    if not submitted:
        return
    title_items = _dedupe(re.split(r"[,;\n]", titles))
    location_items = _dedupe(re.split(r"[,;\n]", locations)) or [""]
    if not title_items:
        st.info("Enter at least one target title.")
        return
    for title in title_items:
        for location in location_items:
            query = " ".join(item for item in (title, location) if item)
            cols = st.columns(3)
            cols[0].link_button("LinkedIn", f"https://www.linkedin.com/jobs/search/?keywords={quote_plus(query)}", width="stretch")
            cols[1].link_button("Indeed", f"https://www.indeed.com/jobs?q={quote_plus(title)}&l={quote_plus(location)}", width="stretch")
            cols[2].link_button("Google Jobs", f"https://www.google.com/search?q={quote_plus(query + ' jobs')}", width="stretch")
            st.caption(query)


def render_method() -> None:
    st.markdown('<div class="hs-eyebrow">How HireSense works</div>', unsafe_allow_html=True)
    st.header("Transparent matching—not a mystery score")
    st.write(
        "HireSense breaks the job description into individual requirements, finds the strongest résumé evidence for "
        "each one, and shows why it is a direct match, related evidence, uncertain, or missing."
    )
    st.markdown(
        """
        | Evidence result | Credit |
        |---|---:|
        | Direct match | 1.00 |
        | Related evidence | 0.65 |
        | Uncertain | 0.25 |
        | Missing | 0.00 |

        Required requirements receive an importance weight of **2** and explicitly preferred requirements receive a
        weight of **1**. Category percentages use the same formula, so the overall score is reproducible from the
        evidence rows—there are no hidden category weights.
        """
    )
    st.latex(
        r"\text{Match Score}=100\times\frac{\sum(\text{evidence credit}\times\text{importance weight})}{\sum\text{importance weight}}"
    )
    st.info("Every conclusion must point back to the uploaded résumé. Unsupported skills, credentials, tools, and metrics are not added automatically.")
    st.warning("HireSense is decision support—not an employer ATS score, hiring decision, or legal/immigration opinion.")


def main() -> None:
    inject_styles()
    initialize_state()
    settings, use_ai = render_sidebar()
    st.markdown(
        """
        <section class="hs-hero">
          <div class="hs-hero-top">
            <div class="hs-brand-chip"><span class="hs-live-dot"></span> HIRESENSE · CAREER INTELLIGENCE</div>
            <div class="hs-privacy-chip">◇ Evidence before claims</div>
          </div>
          <h1><span>HireSense</span> makes every application clearer and stronger.</h1>
          <p class="hs-hero-copy">Upload your résumé, paste a job description, and get a clear match report, priority improvements, and a polished tailored résumé.</p>
          <div class="hs-hero-pills"><span>Understand your fit</span><span>See the proof</span><span>Improve truthfully</span><span>Download and apply</span></div>
        </section>
        <div class="hs-value-grid">
          <div class="hs-value-card"><div class="hs-value-icon">✓</div><strong>Know where you stand</strong><span>See which job requirements your résumé supports and which ones need attention.</span></div>
          <div class="hs-value-card"><div class="hs-value-icon">◎</div><strong>Understand every result</strong><span>Each match is connected to exact evidence from your résumé—not a hidden score.</span></div>
          <div class="hs-value-card"><div class="hs-value-icon">↗</div><strong>Apply with confidence</strong><span>Improve supported content and export a clean, recruiter-ready résumé.</span></div>
        </div>
        <div class="hs-flow">
          <div class="hs-flow-card"><span class="hs-flow-number">01 · ANALYZE</span><strong>Upload and compare</strong><small>Add your résumé and the full job description to create your evidence-based report.</small></div>
          <div class="hs-flow-card"><span class="hs-flow-number">02 · REVIEW</span><strong>Focus on what matters</strong><small>Review your strongest matches, priority gaps, and the résumé evidence behind them.</small></div>
          <div class="hs-flow-card"><span class="hs-flow-number">03 · IMPROVE</span><strong>Tailor and download</strong><small>Create a polished résumé using only experience and skills you already have.</small></div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    analyze_tab, discover_tab, tailor_tab, tracker_tab, method_tab = st.tabs(
        ["1 · Analyze", "2 · Discover", "3 · Tailor & Improve", "4 · Tracker", "5 · How It Works"]
    )
    with analyze_tab:
        render_analyze_input(settings, use_ai)
        render_report(settings, use_ai)
    with discover_tab:
        render_discover()
    with tailor_tab:
        render_tailor_workspace(settings, use_ai)
    with tracker_tab:
        render_tracker()
    with method_tab:
        render_method()
    st.markdown(
        '<div class="hs-footer"><strong>HireSense</strong> · Understand your fit. Improve with evidence. Apply with confidence.</div>',
        unsafe_allow_html=True,
    )


if __name__ == "__main__":
    main()
